"""
Datasheet download regressions.

1. LT Radar datasheet with several tags used to return
       Content-Disposition: attachment; filename=LT_Radar_Datasheet_A,_B.xlsx
   Browsers reject a response whose filename contains an unquoted comma
   (ERR_RESPONSE_HEADERS_MULTIPLE_CONTENT_DISPOSITION), so only single-tag
   downloads worked.
2. An SOP whose datasheet (EG) sheet is hidden could not be saved after the
   other sheets were removed ("The only worksheet of a workbook cannot be hidden").
"""
import io
import json
import re
from types import SimpleNamespace

import openpyxl
import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient

import app.routers.lt_radar as lt
from app.deps import get_current_user

TAGS = ["200-LIT-20201", "200-LIT-20202", "200-FIT-20101", "200-PIT-20105", "200-LIT-20203",
        "200-LIT-20204", "200-LIT-20205", "200-LIT-20206", "200-LIT-20207", "200-LIT-20208",
        "200-LIT-20209", "200-LIT-20210"]


@pytest.fixture()
def client(monkeypatch):
    monkeypatch.setattr(lt, "validate_specs", lambda values: [])   # isolate packaging from spec rules
    app = FastAPI()
    app.include_router(lt.router, prefix="/api")
    app.dependency_overrides[get_current_user] = lambda: SimpleNamespace(id=1, username="u", email="u@x")
    return TestClient(app)


def post(client, tags):
    return client.post("/api/lt-radar/generate",
                       data={"values_json": json.dumps({"tag_no": ", ".join(tags), "make": "X"})})


def browser_safe(content_disposition: str) -> bool:
    """A filename a browser will accept: quoted, ASCII, and no comma anywhere in the header."""
    return ("," not in content_disposition
            and re.search(r'filename="[A-Za-z0-9._-]+\.xlsx"', content_disposition) is not None)


@pytest.mark.parametrize("n", [1, 2, 5, 12])
def test_any_number_of_tags_downloads_with_a_browser_safe_header(client, n):
    r = post(client, TAGS[:n])
    assert r.status_code == 200
    assert browser_safe(r.headers["content-disposition"]), r.headers["content-disposition"]
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert int(r.headers["content-length"]) == len(r.content) > 0
    assert r.headers["x-tag-count"] == str(n)
    wb = openpyxl.load_workbook(io.BytesIO(r.content))      # a real, openable workbook
    text = " ".join(str(c.value) for ws in wb for row in ws.iter_rows() for c in row if c.value)
    for t in TAGS[:n]:
        assert t in text                                      # every selected tag is on the datasheet


def test_single_tag_filename_is_unchanged(client):
    r = post(client, ["200-LIT-20201"])
    assert 'filename="LT_Radar_Datasheet_200-LIT-20201.xlsx"' in r.headers["content-disposition"]


def test_generating_twice_returns_each_selection_correctly(client):
    first = post(client, TAGS[:2])
    second = post(client, TAGS[2:5])
    assert first.headers["x-tag-count"] == "2" and second.headers["x-tag-count"] == "3"
    assert first.headers["content-disposition"] != second.headers["content-disposition"]


def test_hostile_tag_characters_cannot_break_the_header(client):
    r = post(client, ['A"B', "C/D", "E\r\nX-Injected: 1"])
    assert r.status_code == 200
    cd = r.headers["content-disposition"]
    assert "\n" not in cd and "\r" not in cd and browser_safe(cd)


def test_sop_with_hidden_datasheet_sheet_still_generates():
    from utils.sop_datasheet import _generate_xlsx_inplace
    wb = openpyxl.Workbook()
    wb.active.title = "Coversheet"
    eg = wb.create_sheet("EG PIT")
    eg["B2"] = "old"
    eg.sheet_state = "hidden"
    buf = io.BytesIO()
    wb.save(buf)

    out = _generate_xlsx_inplace(buf.getvalue(), "EG PIT", {(1, 1): ("200-LIT-20201", "iodb")})
    result = openpyxl.load_workbook(io.BytesIO(out))
    assert result.sheetnames == ["EG PIT"]
    assert result["EG PIT"]["B2"].value == "200-LIT-20201"
    assert result["EG PIT"].sheet_state == "visible"


# ── SOP (Smart Datasheet Generator) multi-tag packaging, using the repo's sample files ──
ROOT = __import__("pathlib").Path(__file__).resolve().parents[2]
SOP, IODB = ROOT / "Pressure Transmitter.xlsx", ROOT / "IODB.xlsx"


@pytest.mark.skipif(not (SOP.exists() and IODB.exists()), reason="sample SOP/IODB not in repo")
def test_sop_multi_tag_zip_reports_failed_tags_instead_of_blank_sheets():
    import zipfile
    from utils.sop_datasheet import generate_with_report, generate

    args = (SOP.read_bytes(), SOP.name, "DS PIT ", IODB.read_bytes(), "CONTROL VALVE POSTIONER XMTR")
    tags = ["6816-PZI-5004", "6816-PZI-5009", "NO-SUCH-TAG-999", "6816-PZI-5004"]   # 1 missing, 1 duplicate
    data, name, err, failed = generate_with_report(*args, tags, {})

    assert err is None and name == "Datasheets.zip"
    assert failed == [{"tag": "NO-SUCH-TAG-999", "reason": "Tag not found in the IODB"}]
    z = zipfile.ZipFile(io.BytesIO(data))
    assert z.testzip() is None
    assert sorted(z.namelist()) == ["6816-PZI-5004 Datasheet.xlsx", "6816-PZI-5009 Datasheet.xlsx"]
    for n in z.namelist():                                  # each member is a real workbook
        assert openpyxl.load_workbook(io.BytesIO(z.read(n))).sheetnames

    # backward-compatible 3-tuple wrapper keeps working for any other caller
    one = generate(*args, ["6816-PZI-5004"], {})
    assert len(one) == 3 and one[2] is None


@pytest.mark.skipif(not (SOP.exists() and IODB.exists()), reason="sample SOP/IODB not in repo")
def test_sop_all_tags_missing_is_an_error_not_an_empty_success():
    from utils.sop_datasheet import generate_with_report
    data, _name, err, failed = generate_with_report(
        SOP.read_bytes(), SOP.name, "DS PIT ", IODB.read_bytes(),
        "CONTROL VALVE POSTIONER XMTR", ["NOPE-1", "NOPE-2"], {})
    assert data is None and err and "NOPE-1" in err and len(failed) == 2

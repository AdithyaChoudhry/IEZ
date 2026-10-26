"""
I/O Assignment reader.

The layout differs per project, so nothing is positional: the header row is found by scoring rows
for known field aliases, columns are mapped semantically, the detected mapping is returned for the
engineer to confirm or correct, and a correction (field -> column letter) is applied on re-parse.
"""
from __future__ import annotations

import io
from dataclasses import dataclass, field
from typing import Any, Optional

from .common import clean, is_value, norm_header

IO_FIELDS: dict[str, list[str]] = {
    "TAG_NO": ["TAG NO", "TAG NUMBER", "TAG", "INSTRUMENT TAG", "TAG NAME"],
    "LOOP_NO": ["LOOP NO", "LOOP NUMBER", "LOOP"],
    "SERVICE": ["SERVICE"],
    "DESCRIPTION": ["DESCRIPTION", "I O DESCRIPTION", "IO DESCRIPTION", "SIGNAL DESCRIPTION"],
    "STATUS": ["STATUS"],
    "SIGNAL_TYPE_TEXT": ["TYPE OF SIGNAL", "SIGNAL TYPE", "SIGNAL"],
    "WIRE_TYPE": ["WIRE TYPE", "WIRING", "NO OF WIRES"],
    "IO_TYPE": ["IO TYPE", "I O TYPE", "CHANNEL TYPE", "CARD TYPE", "I O"],
    "JB_NO": ["JB NO", "JB TAG", "JB TAG NO", "JUNCTION BOX", "JUNCTION BOX NO", "JUNCTION BOX TAG NO"],
    "JB_TB1": ["JB TB1", "JB TERMINAL 1"],
    "JB_TB2": ["JB TB2", "JB TERMINAL 2"],
    "JB_TB3": ["JB TB3", "JB TERMINAL 3"],
    "JB_PAIR": ["JB PR1", "JB PAIR", "JB PAIR NO"],
    "PLC": ["PLC", "DCS", "PLC DCS", "CONTROLLER", "PLC NO", "SYSTEM"],
    "RACK": ["RACK", "RACK NO", "RACK NUMBER", "RACK NODE", "NODE"],
    "SLOT": ["SLOT", "SLOT NO", "SLOT NUMBER"],
    "MODULE": ["MODULE", "MODULE NO", "MODULE NUMBER", "MODULE TYPE", "CARD", "CARD NO"],
    "CHANNEL": ["CHANNEL", "CHANNEL NO", "CHANNEL NUMBER", "CH NO", "CH"],
    "TERMINAL_STRIP": ["TERMINAL STRIP NUMBER", "TERMINAL STRIP", "TERMINAL STRIP NO", "TERMINAL BLOCK"],
    "TERMINAL": ["TERMINAL", "TERMINAL NO", "TERMINAL NUMBER", "TERMINAL DETAILS"],
    "TERMINAL_CMN": ["TERMINAL NO CMN", "TERMINAL CMN"],
    "TERMINAL_SIGNAL": ["TERMINAL NO SIGNAL", "TERMINAL SIGNAL"],
    "TERMINAL_ZERO": ["TERMINAL NO ZERO", "TERMINAL ZERO"],
    "ADDRESS": ["ADDRESS", "IO ADDRESS", "I O ADDRESS", "PLC ADDRESS"],
    "REMARKS": ["REMARKS", "REMARK"],
}
_ALIAS_TO_FIELD = {a: f for f, aliases in IO_FIELDS.items() for a in aliases}
_HEADER_SCAN_ROWS = 40
_MIN_FIELDS = 3
SPARE_WORDS = {"SPARE", "SPR"}


@dataclass
class IoAssignmentData:
    file_name: str
    rows: list[dict[str, Any]] = field(default_factory=list)
    mapping: dict[str, dict] = field(default_factory=dict)          # field -> {header, column, sheet}
    headers: list[dict] = field(default_factory=list)               # [{sheet, column, header}] for correction dropdowns
    sheets: list[dict] = field(default_factory=list)                # [{sheet, header_row, rows}]
    warnings: list[str] = field(default_factory=list)

    def summary(self) -> dict:
        real = [r for r in self.rows if not r["is_spare"]]
        return {"file_name": self.file_name, "rows": len(self.rows), "tag_rows": len(real),
                "tags": len({r["tag"] for r in real}), "mapping": self.mapping, "headers": self.headers,
                "sheets": self.sheets, "warnings": self.warnings,
                "unmapped_required": [f for f in ("TAG_NO",) if f not in self.mapping]}


def _letter(idx: int) -> str:
    from openpyxl.utils import get_column_letter
    return get_column_letter(idx + 1)


def _score_row(row: tuple) -> dict[str, int]:
    found: dict[str, int] = {}
    for idx, cell in enumerate(row):
        key = norm_header(cell) if cell not in (None, "") else ""
        field_name = _ALIAS_TO_FIELD.get(key)
        if field_name and field_name not in found:
            found[field_name] = idx
    return found


def parse_io_assignment(file_bytes: bytes, file_name: str = "IO_Assignment.xlsx",
                        mapping_override: Optional[dict[str, Optional[str]]] = None) -> IoAssignmentData:
    """
    mapping_override: {field: column letter | None}; applies to every sheet that was detected.
    Raises ValueError (user-facing) if no usable table is found.
    """
    import openpyxl
    from openpyxl.utils import column_index_from_string

    try:
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError(f"The I/O Assignment file could not be opened: {exc}")

    data = IoAssignmentData(file_name=file_name)
    seen_headers: set[tuple[str, str]] = set()
    try:
        for ws in wb.worksheets:
            window = list(ws.iter_rows(min_row=1, max_row=_HEADER_SCAN_ROWS, max_col=80, values_only=True))
            best: Optional[tuple[int, dict[str, int]]] = None
            for i, row in enumerate(window):
                found = _score_row(row)
                if "TAG_NO" in found and len(found) >= _MIN_FIELDS and (best is None or len(found) > len(best[1])):
                    best = (i + 1, found)
            if best is None:
                continue
            header_row, found = best
            header_cells = window[header_row - 1]
            columns = {f: idx for f, idx in found.items()}
            if mapping_override:
                for f, letter in mapping_override.items():
                    if letter in (None, ""):
                        columns.pop(f, None)
                    else:
                        try:
                            columns[f] = column_index_from_string(letter.upper()) - 1
                        except ValueError:
                            data.warnings.append(f"Ignored invalid column '{letter}' for {f}.")
            for idx, h in enumerate(header_cells):
                if h not in (None, "") and (ws.title, _letter(idx)) not in seen_headers:
                    seen_headers.add((ws.title, _letter(idx)))
                    data.headers.append({"sheet": ws.title, "column": _letter(idx), "header": clean(h)})
            for f, idx in columns.items():
                data.mapping.setdefault(f, {"header": clean(header_cells[idx]) if idx < len(header_cells) else "",
                                            "column": _letter(idx), "sheet": ws.title})
            if "TAG_NO" not in columns:
                continue
            count = 0
            max_idx = max(columns.values()) + 1
            for r_idx, row in enumerate(ws.iter_rows(min_row=header_row + 1, max_col=max_idx, values_only=True),
                                        start=header_row + 1):
                vals = {f: clean(row[idx]) if idx < len(row) else "" for f, idx in columns.items()}
                tag = vals.get("TAG_NO", "")
                if not tag:
                    continue
                count += 1
                data.rows.append({"tag": tag, "loop": vals.get("LOOP_NO", ""), "fields": vals, "sheet": ws.title,
                                  "row": r_idx, "is_spare": tag.upper() in SPARE_WORDS or not is_value(tag)})
            data.sheets.append({"sheet": ws.title, "header_row": header_row, "rows": count})
    finally:
        wb.close()

    if not data.rows:
        raise ValueError("No I/O Assignment table found: need a header row with a tag column "
                         "(for example 'Tag Number') and at least two more recognisable columns "
                         "(rack, slot, channel, module, terminal ...).")
    return data


def io_class(io_type: str) -> str:
    """'AI101' -> 'AI', 'DI-03' -> 'DI', 'rtd' -> 'RTD'. '' when the text carries no I/O class."""
    key = norm_header(io_type).replace(" ", "")
    for cls in ("AI", "AO", "DI", "DO", "RTD", "TC", "PI", "FI"):
        if key.startswith(cls):
            return cls
    return ""

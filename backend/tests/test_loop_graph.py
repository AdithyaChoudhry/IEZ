"""Graph designer (Loop Wiring Test): explicit ports/wires, compatibility, validation, renderer, Excel."""
import copy
import io
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))

from openpyxl import load_workbook  # noqa: E402

from utils.loop_test import graph as G  # noqa: E402
from utils.loop_test import graph_excel as GX  # noqa: E402
from utils.loop_test import graph_scene as S  # noqa: E402
from utils.loop_test import graph_templates as T  # noqa: E402
from utils.loop_test import graph_validate as V  # noqa: E402


def _f(v):
    return {"value": v, "source": "IODB", "label": "", "available": True}


def datas(tag="PT-01", jb="JBA-01", rows=1):
    sig = [{"io_type": "AI101", "rack": str(i + 1), "slot": "3", "module": "AI-08", "channel": str(5 + i), "plc": "PLC-01",
            "terminal_strip": "", "terminal": "", "terminal_cmn": "", "terminal_signal": "", "terminal_zero": "", "address": "",
            "description": "", "source": "I/O ASSIGNMENT", "sheet": "S", "row": 7 + i} for i in range(rows)]
    return {tag: {"fields": {"TAG_NO": _f(tag), "INSTRUMENT_TYPE": _f("PRESSURE TRANSMITTER"), "SIGNAL_IO_TYPE": _f("AI"),
                             "JB_TAG": _f(jb), "BRANCH_CABLE_TAG": _f("IC-PT-01"), "BRANCH_CABLE_TYPE": _f("1PX1.5 Sqmm"),
                             "MAIN_CABLE_TAG": _f("MC-JB-01"), "HART": _f("N")}, "signals": sig, "conflicts": []}}


def node(g, typ):
    return next(n for n in g["nodes"] if n["type"] == typ)


def wire(g, a, ap, b, bp):
    e, why = G.add_edge(g, a["id"], ap, b["id"], bp)
    assert e, why
    return e


def design():
    """PT-01 -> JBA-01 -> DCS08+MA -> DCS07+MA (AI card), exactly as the engineer would connect it."""
    d = datas()
    g = T.build_template("AI_2WIRE_STANDARD", "PT-01", "PT-01", d)
    pt, jb, ma = g["nodes"][0], node(g, "junction_box"), node(g, "marshalling_cabinet")
    card = node(g, "plc_ai")
    ma["tag"], card["tag"] = "DCS08+MA", "DCS07+MA"
    for n, t in ((jb, ("15", "16", "S")), (ma, ("15", "16", "S"))):
        for p, v in zip(n["ports"], t):
            p["terminal"] = v
    g["cables"] = [{"id": "c1", "iodb": "branch", "shield": {"type": "individual", "items": [{"id": "S1", "source": "PT-01", "destination": "JB", "termination": "floating", "grounding": ""}]}},
                   {"id": "c2", "iodb": "main", "shield": {"type": "overall", "items": [{"id": "S2", "source": "JB", "destination": "MA", "termination": "control", "grounding": "earth bar"}]}},
                   {"id": "c3", "iodb": "", "tag": "IC-MA-SYS", "type": "1PX1.0", "cores": "2", "shield": {"type": "none", "items": []}}]
    for pp, jp in (("sp", "t1"), ("sn", "t2"), ("sh", "sh")):
        e = wire(g, pt, pp, jb, jp)
        e["cable_id"], e["core"] = "c1", {"sp": "1", "sn": "2", "sh": ""}[pp]
    for jp, mp in (("t1", "t1"), ("t2", "t2"), ("sh", "sh")):
        e = wire(g, jb, jp, ma, mp)
        e["cable_id"], e["core"] = "c2", {"t1": "1", "t2": "2", "sh": ""}[jp]
    for mp, cp in (("t1", "sp"), ("t2", "sn")):
        e = wire(g, ma, mp, card, cp)
        e["cable_id"], e["core"] = "c3", {"t1": "1", "t2": "2"}[mp]
    return g, d


def test_template_creates_blocks_and_ports_but_no_wires():
    g = T.build_template("AI_2WIRE_STANDARD", "PT-01", "PT-01", datas())
    assert [n["type"] for n in g["nodes"]] == ["pressure_transmitter", "junction_box", "marshalling_cabinet", "plc_ai", "control_function"]
    assert g["edges"] == []                               # nothing is inferred
    assert [p["label"] for p in g["nodes"][0]["ports"]] == ["SIG+", "SIG-", "SHIELD"]
    assert [n["zone"] for n in g["nodes"]] == ["field", "jb", "marshalling", "system", "control"]


def test_four_wire_has_separate_power_ports_and_a_power_supply():
    g = T.build_template("AI_4WIRE_STANDARD", "PT-01", "PT-01", datas())
    dev = g["nodes"][0]
    assert [p["kind"] for p in dev["ports"]] == ["power", "power", "signal", "signal", "shield"]
    assert any(n["type"] == "power_supply" for n in g["nodes"])


def test_acceptance_designed_loop_validates_and_draws_exactly_the_configured_wires():
    g, d = design()
    check = V.validate_graph(g, d)
    errs = [c for c in check["checks"] if c["status"] == "ERROR"]
    assert not errs, errs
    scene = S.build_scene(S.normalize_positions(g, d), d)
    assert set(scene["edges"]) == {e["id"] for e in g["edges"]} and len(g["edges"]) == 8
    texts = [i["s"] for i in scene["items"] if i["t"] == "text"]
    assert any("INSTRUMENT LOOP DIAGRAM FOR PRESSURE TRANSMITTER PT-01" == t for t in texts)
    assert "15  +" in texts and "16  -" in texts and "DCS08+MA" in texts and "DCS07+MA" in texts


def test_moving_blocks_moves_the_wires_but_never_changes_connections():
    g, d = design()
    before = [(e["id"], e["from_node"], e["from_port"], e["to_node"], e["to_port"]) for e in g["edges"]]
    jb = node(g, "junction_box")
    ends_before = S.build_scene(g, d)["edges"]
    jb["x"] += 40
    jb["y"] += 130
    after_scene = S.build_scene(g, d)
    assert [(e["id"], e["from_node"], e["from_port"], e["to_node"], e["to_port"]) for e in g["edges"]] == before
    moved = [eid for eid, e in after_scene["edges"].items() if e["pts"] != ends_before[eid]["pts"]]
    touching = {e["id"] for e in g["edges"] if jb["id"] in (e["from_node"], e["to_node"])}
    assert set(moved) == touching                          # only wires attached to the moved block follow it


def test_change_jb_terminal_and_plc_channel_changes_only_that_endpoint():
    g, d = design()
    jb = node(g, "junction_box")
    base = copy.deepcopy(g["edges"])
    next(p for p in jb["ports"] if p["id"] == "t1")["terminal"] = "17"
    assert g["edges"] == base
    texts = [i["s"] for i in S.build_scene(g, d)["items"] if i["t"] == "text"]
    assert "17  +" in texts and "15  +" in texts            # JB shows 17, marshalling still 15
    d2 = datas(rows=2)
    card = node(g, "plc_ai")
    card["props"]["io_index"] = 0
    t0 = [i["s"] for i in S.build_scene(g, d2)["items"] if i["t"] == "text"]
    card["props"]["io_index"] = 1
    t1 = [i["s"] for i in S.build_scene(g, d2)["items"] if i["t"] == "text"]
    assert "Channel: 5" in t0 and "Channel: 6" in t1 and g["edges"] == base


def test_incompatible_connections_are_refused_with_a_reason():
    g, d = design()
    do = G.new_node("plc_do", "n99")
    g["nodes"].append(do)
    pt = g["nodes"][0]
    shield_to_sig = G.can_connect(g, pt["id"], "sh", node(g, "plc_ai")["id"], "sp")
    assert not shield_to_sig[0] and "shield" in shield_to_sig[1].lower()
    ok, why = G.can_connect(g, node(g, "marshalling_cabinet")["id"], "t1", "n99", "sp")
    assert not ok                                           # port already used / or family mismatch is reported
    g2 = T.build_template("AI_4WIRE_STANDARD", "PT-01", "PT-01", d)
    dev, ps, card = g2["nodes"][0], node(g2, "power_supply"), node(g2, "plc_ai")
    ok, why = G.can_connect(g2, ps["id"], "pp", card["id"], "sp")
    assert not ok and "power" in why.lower()                # 24 V cannot go straight onto an AI signal port
    ok, why = G.can_connect(g2, dev["id"], "pp", ps["id"], "pp")
    assert ok, why
    ok, why = G.can_connect(g2, dev["id"], "pp", ps["id"], "pn")
    assert not ok and "polarity" in why.lower()
    g3 = T.build_template("DO_SOLENOID", "PT-01", "PT-01", d)
    sol = node(g3, "solenoid")
    ok, why = G.can_connect(g3, sol["id"], "sp", node(g3, "plc_do")["id"], "sp")
    assert ok, why
    g4 = T.build_template("AI_2WIRE_STANDARD", "PT-01", "PT-01", d)
    g4["nodes"].append(G.new_node("plc_di", "n98"))
    ok, why = G.can_connect(g4, g4["nodes"][0]["id"], "sp", "n98", "sp")
    assert not ok and "AI" in why and "DI" in why


def test_validation_catches_duplicate_cores_missing_terminals_and_unconnected_ports():
    g, d = design()
    g["edges"][1]["core"] = g["edges"][0]["core"]           # two wires on core 1 of the same cable
    jb = node(g, "junction_box")
    next(p for p in jb["ports"] if p["id"] == "t2")["terminal"] = ""
    c = {x["id"]: x for x in V.validate_graph(g, d)["checks"]}
    assert c["dup_cores"]["status"] == "ERROR" and c["terminals"]["status"] == "ERROR"
    g2 = T.build_template("AI_2WIRE_STANDARD", "PT-01", "PT-01", d)
    c2 = {x["id"]: x for x in V.validate_graph(g2, d)["checks"]}
    assert c2["required_ports"]["status"] == "ERROR" and c2["paths"]["status"] == "PASS"


def test_jb_and_cable_must_match_iodb_unless_the_engineer_accepts_the_difference():
    g, d = design()
    jb = node(g, "junction_box")
    jb["props"]["jb_tag"] = "JBA-99"
    assert {x["id"]: x for x in V.validate_graph(g, d)["checks"]}["jb_match"]["status"] == "ERROR"
    jb["props"]["accept_difference"] = True
    assert {x["id"]: x for x in V.validate_graph(g, d)["checks"]}["jb_match"]["status"] == "PASS"


def test_motor_vfd_each_signal_is_its_own_connection_path():
    d = datas("M-01")
    g = T.build_template("MOTOR_VFD", "M-01", "M-01", d)
    assert [n["type"] for n in g["nodes"]] == ["motor", "vfd", "plc", "control_function"]
    assert not [p for n in g["nodes"] for p in n["ports"] if p.get("signal_id")]     # nothing pre-selected
    T.apply_signals(g, ["start", "speed_reference", "run_feedback"])
    vfd, plc = node(g, "vfd"), node(g, "plc")
    assert {p["sig"] for p in vfd["ports"] if p.get("signal_id")} == {"DO", "AO", "DI"}
    e = wire(g, plc, "start_p", vfd, "start_p")
    assert e["kind"] == "signal"
    ok, why = G.can_connect(g, plc["id"], "start_n", vfd["id"], "run_feedback_n")
    assert not ok                                           # DO command cannot land on a DI feedback port
    T.apply_signals(g, ["speed_reference"])                 # deselecting removes that signal's wires
    assert g["edges"] == [] and not any(p["id"].startswith("start") for p in plc["ports"])


def test_duplicate_loop_rereads_source_values_and_clears_loop_specific_ones():
    g, d = design()
    d2 = datas("PT-02", jb="JBA-02")
    g2, report = T.duplicate_for_tag(g, "PT-02", "PT-02", d2)
    jb = node(g2, "junction_box")
    assert g2["tag"] == "PT-02" and jb["props"]["jb_tag"] == "JBA-02"
    assert all(not p["terminal"] for p in jb["ports"])
    assert len(g2["edges"]) == len(g["edges"])              # the wiring structure is kept
    assert next(c for c in g2["cables"] if c["id"] == "c3")["tag"] == ""
    assert any("terminal" in r for r in report)


def test_png_and_excel_come_from_the_same_graph():
    g, d = design()
    S.normalize_positions(g, d)
    scene = S.build_scene(g, d, {"project": "P1"})
    png = S.render_png(scene)
    assert png[:8] == b"\x89PNG\r\n\x1a\n"
    out = GX.build_workbook([{"tag": "PT-01", "graph": g, "datas": d, "scene": scene, "validation": V.validate_graph(g, d)}])
    ws = load_workbook(io.BytesIO(out))["PT-01"]
    assert len(ws._images) == 1
    wires = [r for r in ws.iter_rows(values_only=True) if r[0] and str(r[0]).startswith("W") and r[0] != "Wire"]
    assert len(wires) == len(g["edges"])

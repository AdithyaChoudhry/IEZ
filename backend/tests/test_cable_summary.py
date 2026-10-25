"""
CABLE SUMMARY ("CABLE QUANTITIES") sheet of the Cable Schedule workbook.
Numbered comments map to the TEST cases in the feature specification.
"""
import io
import time
from pathlib import Path

import openpyxl
import pandas as pd
import pytest

from utils import cable_summary as cs
from utils.cable_summary import (CABLE_SUMMARY_SHEET, append_cable_summary_sheet, build_cable_summary,
                                 parse_length)

ROOT = Path(__file__).resolve().parents[2]
TEMPLATE = ROOT / "Cable_schedule.xlsx"
BT, BL, MT, ML = "INSTRUMENT TO JB CABLE TYPE", "BRANCH CABLE LENGTH", "MAIN CABLE TYPE", "MAIN CABLE LENGTH(MTRS)"


def frame(rows):
    """rows: (tag, branch_type, branch_len, main_type, main_len)"""
    return pd.DataFrame(rows, columns=["TAG NO", BT, BL, MT, ML])


def rows_of(result):
    return [(r["cable_description"], r["quantity"]) for r in result.rows]


# ── spec example ────────────────────────────────────────────────────────────
def test_specification_example():
    df = frame([("TAG 1", "1Px1.5Sq.mm", 27, "2Px1.5Sq.mm", 50),
                ("TAG 2", "1Px1.5Sq.mm", 35, "2Px1.5Sq.mm", 40),
                ("TAG 3", "6Px1.5Sq.mm", 397, "2Px1.5Sq.mm", 20)])
    assert rows_of(build_cable_summary(df)) == [("1Px1.5Sq.mm", 62), ("2Px1.5Sq.mm", 110), ("6Px1.5Sq.mm", 397)]
    assert all(r["units"] == "METERS" for r in build_cable_summary(df).rows)


def test_first_example_in_section_3():
    df = frame([("TAG 1", "1Px1.5Sq.mm", 27, "2Px1.5Sq.mm", 35), ("TAG 2", "1Px1.5Sq.mm", 32, "2Px1.5Sq.mm", 40)])
    assert rows_of(build_cable_summary(df)) == [("1Px1.5Sq.mm", 59), ("2Px1.5Sq.mm", 75)]


# ── 1-5 grouping and sources ────────────────────────────────────────────────
def test_1_branch_only():
    assert rows_of(build_cable_summary(frame([("A", "1Px1.5Sq.mm", 10, None, None)]))) == [("1Px1.5Sq.mm", 10)]


def test_2_main_only():
    assert rows_of(build_cable_summary(frame([("A", None, None, "2Px2.5Sq.mm", 80)]))) == [("2Px2.5Sq.mm", 80)]


def test_3_same_type_in_branch_and_main_is_one_row():
    df = frame([("A", "1Px1.5Sq.mm", 100, "1Px1.5Sq.mm", 250)])
    assert rows_of(build_cable_summary(df)) == [("1Px1.5Sq.mm", 350)]


def test_4_same_type_in_many_rows_is_summed():
    df = frame([(f"T{i}", "1Px1.5Sq.mm", 10, None, None) for i in range(6)])
    assert rows_of(build_cable_summary(df)) == [("1Px1.5Sq.mm", 60)]


def test_5_different_types_are_separate_rows_sorted_naturally():
    df = frame([("A", "12Px1.5Sq.mm", 5, "6Px1.5Sq.mm", 1), ("B", "1Px2.5Sq.mm", 2, "2Px1.5Sq.mm", 3),
                ("C", "1Px1.5Sq.mm", 4, "4Px1.5Sq.mm", 6)])
    assert [d for d, _ in rows_of(build_cable_summary(df))] == [
        "1Px1.5Sq.mm", "1Px2.5Sq.mm", "2Px1.5Sq.mm", "4Px1.5Sq.mm", "6Px1.5Sq.mm", "12Px1.5Sq.mm"]


def test_no_double_counting_between_branch_and_main_lengths():
    df = frame([("A", "1Px1.5Sq.mm", 27, "2Px1.5Sq.mm", 35)])
    assert rows_of(build_cable_summary(df)) == [("1Px1.5Sq.mm", 27), ("2Px1.5Sq.mm", 35)]


# ── 6-8 blanks, invalid, decimals ───────────────────────────────────────────
def test_6_blank_cable_type_is_ignored():
    df = frame([("A", None, 50, "", 60), ("B", "  ", 5, "-", 7), ("C", "TBA", 5, "NA", 7)])
    r = build_cable_summary(df)
    assert r.rows == [] and r.counts["blank_length"] == 0       # nothing to flag: no cable at all


def test_7_blank_or_non_numeric_length_is_skipped_and_flagged():
    df = frame([("A", "1Px1.5Sq.mm", None, "2Px1.5Sq.mm", "27 m"),
                ("B", "1Px1.5Sq.mm", "abc", "2Px1.5Sq.mm", ""),
                ("C", "1Px1.5Sq.mm", 10, "2Px1.5Sq.mm", 20)])
    r = build_cable_summary(df, "TAG NO")
    assert rows_of(r) == [("1Px1.5Sq.mm", 10), ("2Px1.5Sq.mm", 20)]    # never guessed, never zero rows
    assert r.counts["blank_length"] == 2 and r.counts["invalid_length"] == 2
    assert any("'27 m'" in d and "tag A" in d for d in r.details)
    assert any("non-numeric" in w for w in r.warnings)


def test_8_decimals_and_numeric_strings_sum_exactly():
    df = frame([("A", "1Px1.5Sq.mm", 0.1, None, None), ("B", "1Px1.5Sq.mm", "0.2", None, None),
                ("C", "1Px1.5Sq.mm", 397.5, None, None), ("D", "2Px1.5Sq.mm", 1.0, None, None)])
    assert rows_of(build_cable_summary(df)) == [("1Px1.5Sq.mm", 397.8), ("2Px1.5Sq.mm", 1)]
    q = build_cable_summary(df).rows[1]["quantity"]
    assert isinstance(q, int)                                            # whole numbers are shown as 1, not 1.0


def test_negative_and_zero_lengths():
    df = frame([("A", "1Px1.5Sq.mm", -5, "2Px1.5Sq.mm", 0), ("B", "1Px1.5Sq.mm", 4, None, None)])
    r = build_cable_summary(df)
    assert rows_of(r) == [("1Px1.5Sq.mm", 4)] and r.counts["invalid_length"] == 1 and r.counts["zero_length"] == 1


def test_parse_length_cases():
    assert parse_length(27)[1] == "ok" and parse_length("27.5")[0] == cs.Decimal("27.5")
    assert parse_length(None)[1] == "blank" and parse_length(float("nan"))[1] == "blank"
    assert parse_length("27 m")[1] == "invalid" and parse_length(True)[1] == "invalid"
    assert parse_length(float("inf"))[1] == "invalid"


# ── matching rules ──────────────────────────────────────────────────────────
def test_exact_matching_after_whitespace_normalisation_only():
    df = frame([("A", "1Px1.5Sq.mm", 10, None, None), ("B", " 1Px1.5Sq.mm ", 5, None, None),
                ("C", "1Px1.5Sq.mm ", 1, None, None), ("D", "1P x 1.5 Sq.mm", 100, None, None)])
    r = build_cable_summary(df)
    assert rows_of(r) == [("1P x 1.5 Sq.mm", 100), ("1Px1.5Sq.mm", 16)]    # whitespace runs merged; look-alike NOT merged
    assert any("look alike" in w for w in r.warnings)


def test_case_differences_are_not_merged():
    df = frame([("A", "1Px1.5Sq.mm", 1, None, None), ("B", "1px1.5sq.mm", 2, None, None)])
    assert len(build_cable_summary(df).rows) == 2


def test_missing_source_columns_are_reported_not_fatal():
    df = pd.DataFrame({"TAG NO": ["A"], BT: ["1Px1.5Sq.mm"], BL: [10]})        # no main cable columns
    r = build_cable_summary(df)
    assert rows_of(r) == [("1Px1.5Sq.mm", 10)] and MT in r.missing_columns
    assert build_cable_summary(pd.DataFrame({"TAG NO": ["A"]})).rows == []


def test_header_spelling_tolerance():
    df = pd.DataFrame({"Instrument to JB Cable Type ": ["1Px1.5Sq.mm"], "branch  cable length": [3],
                       "MAIN CABLE TYPE": ["2Px1.5Sq.mm"], "MAIN CABLE LENGTH (MTRS)": [4]})
    assert rows_of(build_cable_summary(df)) == [("1Px1.5Sq.mm", 3), ("2Px1.5Sq.mm", 4)]


def test_9_large_project_is_fast():
    rows = [(f"T{i}", f"{1 + i % 6}Px1.5Sq.mm", 10 + i % 7, f"{2 + i % 4}Px2.5Sq.mm", 20 + i % 5) for i in range(5000)]
    t = time.time()
    r = build_cable_summary(frame(rows), "TAG NO")
    assert time.time() - t < 1.0 and len(r.rows) >= 8


# ── worksheet ───────────────────────────────────────────────────────────────
def test_new_summary_sheet_is_appended_last_when_template_has_none():
    wb = openpyxl.Workbook()
    wb.active.title = "Cable Schedule"
    wb.create_sheet("Other")
    r = build_cable_summary(frame([("A", "1Px1.5Sq.mm", 27, "2Px1.5Sq.mm", 50)]))
    append_cable_summary_sheet(wb, r)
    assert wb.sheetnames == ["Cable Schedule", "Other", "CABLE SUMMARY"]
    ws = wb["CABLE SUMMARY"]
    assert ws["A1"].value == "CABLE QUANTITIES"
    assert [c.value for c in ws[3][:4]] == ["SL.NO", "CABLE DESCRIPTION", "UNITS", "QUANTITY"]
    assert [c.value for c in ws[4][:4]] == [1, "1Px1.5Sq.mm", "METERS", 27]


def test_empty_summary_has_headers_and_a_note_but_no_zero_row():
    wb = openpyxl.Workbook()
    append_cable_summary_sheet(wb, build_cable_summary(frame([("A", None, None, None, None)])))
    ws = wb["CABLE SUMMARY"]
    assert wb.sheetnames[-1] == "CABLE SUMMARY"
    assert "No cable type" in ws["B4"].value and ws["D4"].value is None
    assert not any(c.value == 0 for row in ws.iter_rows() for c in row)


# ── full Cable Schedule generation (real template) ──────────────────────────
def iodb_bytes(rows, with_cable_columns=True):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "IODB"
    head = ["TAG NO", "JUNCTION BOX"] + ([BT, BL, MT, ML] if with_cable_columns else [])
    ws.append(head)
    for i, (tag, bt, bl, mt, ml) in enumerate(rows):
        ws.append([tag, f"JB-{1 + i % 2}"] + ([bt, bl, mt, ml] if with_cable_columns else []))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


SPEC_ROWS = [("TAG 1", "1Px1.5Sq.mm", 27, "2Px1.5Sq.mm", 50), ("TAG 2", "1Px1.5Sq.mm", 35, "2Px1.5Sq.mm", 40),
             ("TAG 3", "6Px1.5Sq.mm", 397, "2Px1.5Sq.mm", 20)]


@pytest.fixture(scope="module")
def generated():
    if not TEMPLATE.exists():
        pytest.skip("Cable_schedule.xlsx template not in repo")
    from utils.template_processor import process_cable_schedule
    tmpl = TEMPLATE.read_bytes()
    out, name, err = process_cable_schedule(io.BytesIO(iodb_bytes(SPEC_ROWS)), io.BytesIO(tmpl), "JUNCTION BOX", "TAG NO")
    assert err is None and name == "Cable_Schedule.xlsx"
    return openpyxl.load_workbook(io.BytesIO(out)), tmpl


def test_10_cable_summary_is_last_sheet_with_expected_content(generated):
    wb, _ = generated
    assert wb.sheetnames[-1] == CABLE_SUMMARY_SHEET and wb.sheetnames.count(CABLE_SUMMARY_SHEET) == 1
    ws = wb[CABLE_SUMMARY_SHEET]
    assert ws.sheet_state == "visible"
    assert ws["A6"].value == "CABLE QUANTITIES"
    assert [c.value for c in ws[7][:4]] == ["SL.NO", "CABLE DESCRIPTION", "UNITS", "QUANTITY"]
    data = [[c.value for c in ws[r][:4]] for r in (8, 9, 10)]
    assert data == [[1, "1Px1.5Sq.mm", "METERS", 62], [2, "2Px1.5Sq.mm", "METERS", 110], [3, "6Px1.5Sq.mm", "METERS", 397]]
    assert ws["C11"].value == "TOTAL" and ws["D11"].value == "=ROUNDUP(SUM(D8:D10),0)"   # template's own total convention
    assert ws["B12"].value is None                                                          # no stale rows below
    assert ws.print_area.endswith("$D$11")


def test_template_formatting_is_reused(generated):
    ws = generated[0][CABLE_SUMMARY_SHEET]
    assert ws["A7"].fill.fgColor.rgb.endswith("CCFFFF") and ws["A7"].font.b           # template header style
    assert ws["B9"].font.name == "Calibri" and ws["D9"].alignment.horizontal == "center"
    assert ws.column_dimensions["B"].width > 50                                       # template column widths kept
    assert len(ws._images) >= 0


def test_existing_sheets_unchanged_and_in_same_relative_order(generated, monkeypatch):
    wb, tmpl = generated
    from utils.template_processor import process_cable_schedule
    monkeypatch.setattr(cs, "append_cable_summary_sheet", lambda *a, **k: None)         # baseline: feature off
    out, _, err = process_cable_schedule(io.BytesIO(iodb_bytes(SPEC_ROWS)), io.BytesIO(tmpl), "JUNCTION BOX", "TAG NO")
    assert err is None
    base = openpyxl.load_workbook(io.BytesIO(out))
    assert [s for s in wb.sheetnames if s != CABLE_SUMMARY_SHEET] == [s for s in base.sheetnames if s != CABLE_SUMMARY_SHEET]
    for name in base.sheetnames:
        if name == CABLE_SUMMARY_SHEET:
            continue
        a, b = base[name], wb[name]
        assert a.sheet_state == b.sheet_state and a.max_row == b.max_row
        assert [[c.value for c in r] for r in a.iter_rows()] == [[c.value for c in r] for r in b.iter_rows()], name


def test_iodb_without_cable_columns_still_produces_schedule_and_empty_summary():
    if not TEMPLATE.exists():
        pytest.skip("template missing")
    from utils.template_processor import process_cable_schedule
    out, _, err = process_cable_schedule(io.BytesIO(iodb_bytes(SPEC_ROWS, with_cable_columns=False)),
                                         io.BytesIO(TEMPLATE.read_bytes()), "JUNCTION BOX", "TAG NO")
    assert err is None
    wb = openpyxl.load_workbook(io.BytesIO(out))
    assert wb.sheetnames[-1] == CABLE_SUMMARY_SHEET
    assert "Cable columns not found" in str(wb[CABLE_SUMMARY_SHEET]["B8"].value)


def test_summary_failure_never_blocks_the_schedule(monkeypatch):
    if not TEMPLATE.exists():
        pytest.skip("template missing")
    from utils.template_processor import process_cable_schedule
    monkeypatch.setattr(cs, "build_cable_summary", lambda *a, **k: (_ for _ in ()).throw(RuntimeError("boom")))
    out, name, err = process_cable_schedule(io.BytesIO(iodb_bytes(SPEC_ROWS)), io.BytesIO(TEMPLATE.read_bytes()),
                                            "JUNCTION BOX", "TAG NO")
    assert err is None and out and name == "Cable_Schedule.xlsx"


def test_9b_summary_adds_little_time_to_a_large_generation():
    if not TEMPLATE.exists():
        pytest.skip("template missing")
    from utils.template_processor import process_cable_schedule
    rows = [(f"T{i:04d}", f"{1 + i % 6}Px1.5Sq.mm", 10 + i % 7, f"{2 + i % 4}Px2.5Sq.mm", 20 + i % 5) for i in range(400)]
    t = time.time()
    out, _, err = process_cable_schedule(io.BytesIO(iodb_bytes(rows)), io.BytesIO(TEMPLATE.read_bytes()),
                                         "JUNCTION BOX", "TAG NO")
    assert err is None
    ws = openpyxl.load_workbook(io.BytesIO(out))[CABLE_SUMMARY_SHEET]
    assert ws["D8"].value is not None and time.time() - t < 120
    summary_time = time.time()
    build_cable_summary(frame(rows), "TAG NO")
    assert time.time() - summary_time < 0.5

"""Excel output for graph loops: the drawing is the SAME scene the preview shows (rendered to an image), followed by
tables generated from the same graph - wire / connection schedule, cable schedule, block list, validation notes."""
from __future__ import annotations

import io
from typing import Any

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from .excel_out import safe_sheet_name
from .graph import find_node, find_port
from .graph_scene import LIBRARY, build_scene, node_view, render_png, resolve_cable

THIN = Side(style="thin", color="000000")
F_BASE = Font(name="Arial", size=9)
F_BOLD = Font(name="Arial", size=9, bold=True)
F_HEAD = Font(name="Arial", size=12, bold=True)
F_MISS = Font(name="Arial", size=9, italic=True, color="C00000")
FILL = PatternFill("solid", fgColor="E6E6E6")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _table(ws, r: int, title: str, head: list[str], rows: list[list[Any]]) -> int:
    ws.cell(r, 1, title).font = F_HEAD
    r += 1
    for c, h in enumerate(head, 1):
        cell = ws.cell(r, c, h)
        cell.font, cell.fill, cell.border = F_BOLD, FILL, BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in rows:
        r += 1
        for c, v in enumerate(row, 1):
            cell = ws.cell(r, c, v)
            miss = isinstance(v, str) and ("MISSING" in v or v in ("?", "NOT DEFINED"))
            cell.font, cell.border = (F_MISS if miss else F_BASE), BORDER
            cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=False)
    return r + 2


def port_ref(node: dict, port_id: str) -> tuple[str, str]:
    p = find_port(node, port_id) or {}
    return (p.get("terminal") or "—"), p.get("label", port_id)


def build_workbook(items: list[dict]) -> bytes:
    """items: [{tag, graph, datas, scene, validation, meta}]"""
    wb = Workbook()
    wb.remove(wb.active)
    used: set[str] = set()
    for it in items:
        g, datas, scene, val = it["graph"], it["datas"], it["scene"], it["validation"]
        ws = wb.create_sheet(safe_sheet_name(it["tag"], used))
        ws.sheet_view.showGridLines = False
        for i, w in enumerate((9, 22, 14, 12, 22, 14, 12, 18, 18, 8, 14, 9, 22, 30), 1):
            ws.column_dimensions[chr(64 + i)].width = w
        png = render_png(scene, scale=1.5)
        img = XLImage(io.BytesIO(png))
        target_w = 1500
        ratio = target_w / img.width
        img.width, img.height = int(img.width * ratio), int(img.height * ratio)
        ws.add_image(img, "A2")
        r = 2 + int(img.height / 20) + 2
        nodes = {n["id"]: n for n in g["nodes"]}
        cab = {c["id"]: resolve_cable(c, g, datas) for c in g.get("cables", [])}
        rows = []
        for e in g["edges"]:
            na, nb = nodes[e["from_node"]], nodes[e["to_node"]]
            ta, la = port_ref(na, e["from_port"])
            tb, lb = port_ref(nb, e["to_port"])
            c = cab.get(e.get("cable_id", ""), {})
            sh = (c.get("shield") or {}).get("type", "") if c else ""
            rows.append([e["wire_no"], node_view(na, g, datas)["title"], la, ta, node_view(nb, g, datas)["title"], lb, tb,
                         c.get("tag") or ("—" if not c else "MISSING"), c.get("type") or "—", e.get("core") or "—", e.get("signal", ""),
                         e["kind"], {"": "NOT DEFINED" if c else "—"}.get(sh, sh) if c else "—", e.get("remarks", "")])
        r = _table(ws, r, "CONNECTION (WIRE) SCHEDULE - exactly the wires configured on the canvas",
                   ["Wire", "From block", "From port", "From term.", "To block", "To port", "To term.", "Cable tag", "Cable type", "Core",
                    "Signal", "Kind", "Shield", "Remarks"], rows)
        crow = []
        for cid, c in cab.items():
            sh = c["shield"] or {}
            crow.append([cid, c["tag"] or "MISSING", c["type"] or "MISSING", c["length"] or "MISSING", c["route"] or "—", c["cores"] or "—", c["pairs"] or "—",
                         sh.get("type") or "NOT DEFINED",
                         "; ".join(f"{i.get('id', '')}: {i.get('source', '')}→{i.get('destination', '')} ({i.get('termination', '')}, earth {i.get('grounding', '') or '—'})"
                                   for i in sh.get("items", [])) or "—"])
        r = _table(ws, r, "CABLE SCHEDULE", ["Id", "Cable tag", "Type", "Length (m)", "Route", "Cores", "Pairs", "Shield", "Shield details"], crow)
        brow = []
        for n in g["nodes"]:
            v = node_view(n, g, datas)
            for p in n["ports"]:
                brow.append([v["heading"].title(), v["title"], p["label"], p.get("terminal") or "—", p["kind"], p["pol"] or "—",
                             sum(1 for e in g["edges"] if (e["from_node"], e["from_port"]) == (n["id"], p["id"]) or (e["to_node"], e["to_port"]) == (n["id"], p["id"]))])
        r = _table(ws, r, "BLOCKS AND PORTS", ["Block", "Tag", "Port", "Terminal", "Kind", "Pol.", "Wires"], brow)
        notes = [[c["status"], c["label"], c["detail"]] for c in val["checks"] if c["status"] != "PASS"]
        if notes:
            r = _table(ws, r, "VALIDATION NOTES (generated with these warnings)", ["Status", "Check", "Detail"], notes)
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

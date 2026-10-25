"""
Master Document List (MDL) integration for the Coversheet Generator.

Numbered comments map to the TEST cases in the feature specification.
"""
import io
from pathlib import Path
from types import SimpleNamespace

import openpyxl
import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient

import app.database as database
from app.deps import get_current_user, get_db
from app.routers.coversheet import router as coversheet_router
from app.routers.coversheet_mdl import router as mdl_router
from sqlalchemy.orm import sessionmaker
from utils import mdl_parser
from utils.mdl_parser import MdlFormatError, parse_mdl, split_document_number

REAL_MDL = Path("/Users/adithyachoudhrym/Downloads/Input Files/PC286-10P195-A0001-001_R4_MDL.xlsx")
PID, OTHER = "10P211", "10P204"
NAME = "PC286-10P195-A0001-001_R4_MDL.xlsx"

ROWS = [  # (S.No, discipline, doc no, description, category, remarks)
    (1, "Project", "PC286-10P195-A0001-001", "Master Documents List", "A", ""),
    (None, "INSTRUMENTATION", None, None, None, None),                       # section heading
    (2, "Instrumentation", "PC286-10P195-I0003-100", "Instrument Datasheets for New WWTP", "A", ""),
    (3, "Instrumentation", "PC286-10P195-I0003-200", "Instrument Datasheets for Storm Water Guard Pond", "A", ""),
    (4, "Instrumentation", "PC286-10P195-I0003-300", "Instrument Datasheets for ZLD Unit", "A", ""),
    ("#REF!", "Mechanical", "PC286-10P195-D0032-301", "Fabrication drawings of UF Feed Collection Tank", "A", ""),
    (7, "Mechanical", "PC286-10P195-D0032-301", "Fabrication drawings of RO Feed Tank", "A", ""),   # duplicate DOC NO
    (8, "Process", "PC286-10P195-D0015-201", "P&ID for Storm Water Guard Pond", "I", ""),
    (9, "Process", None, "Row without a document number", "A", ""),           # not a record
    (10, "Process", "PC286-10P195-X0001-001", None, "A", ""),                 # not a record
]


def make_mdl(rows=ROWS, revision=4, columns=None, sheet="MDL", add_cover=True, header_row=3) -> bytes:
    columns = columns or ["S.No", "DISCIPLINE", "DOC NO", "DESCRIPTION", "Doc Category", "Remarks"]
    wb = openpyxl.Workbook()
    cover = wb.active
    cover.title = "Cover"
    if add_cover:
        cover["B19"], cover["D19"] = "DOC. NUMBER", "PC286-10P195-A0001-001"      # must NOT become a record
        cover["K19"], cover["L19"] = "REV", revision
    ws = wb.create_sheet(sheet)
    ws["A1"] = "MASTER DOCUMENTS LIST FOR PROJECT"
    for c, name in enumerate(columns, 1):
        ws.cell(header_row, c, name)
    for r, row in enumerate(rows, header_row + 1):
        order = {"S.No": 0, "DISCIPLINE": 1, "DOC NO": 2, "DESCRIPTION": 3, "Doc Category": 4, "Remarks": 5}
        for c, name in enumerate(columns, 1):
            ws.cell(r, c, row[order.get(name, c - 1)])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.fixture()
def client(engine, db, monkeypatch):
    monkeypatch.setattr(database, "SessionLocal", sessionmaker(bind=engine))
    app = FastAPI()
    for r in (mdl_router, coversheet_router):
        app.include_router(r, prefix="/api")
    app.dependency_overrides[get_current_user] = lambda: SimpleNamespace(id=3, username="alice", email="a@x")
    app.dependency_overrides[get_db] = lambda: db
    return TestClient(app)


def upload(client, data=None, project=PID, name=NAME, selected=None):
    form = {"selected_record_id": selected} if selected else {}
    return client.post(f"/api/projects/{project}/coversheet/mdl/upload", data=form,
                       files={"file": (name, data if data is not None else make_mdl())})


def search(client, q="", project=PID, **kw):
    return client.get(f"/api/projects/{project}/coversheet/mdl/search", params={"q": q, **kw}).json()


# ── 1 / 12: parsing, structure validation ───────────────────────────────────
def test_parser_reads_records_ignores_sections_and_cover_sheet():
    r = parse_mdl(make_mdl(), NAME)
    assert r.source_sheet == "MDL" and r.header_row == 3
    nums = [x.document_number for x in r.records]
    assert len(r.records) == 7 and r.section_rows_ignored == 3      # heading + the 2 half-filled rows
    assert "PC286-10P195-A0001-001" in nums and nums.count("PC286-10P195-A0001-001") == 1   # Cover sheet not used
    assert all(x.source_file == NAME and x.source_sheet == "MDL" for x in r.records)
    assert {x.s_no for x in r.records} >= {"1", "2"} and "" in {x.s_no for x in r.records}   # '#REF!' S.No tolerated


def test_header_columns_found_by_name_not_position():
    shuffled = ["Remarks", "Doc Category", "DESCRIPTION", "DOC NO", "DISCIPLINE", "S.No"]
    r = parse_mdl(make_mdl(columns=shuffled), NAME)
    assert len(r.records) == 7
    assert r.records[1].document_number == "PC286-10P195-I0003-100"
    assert r.records[1].document_name == "Instrument Datasheets for New WWTP"


def test_header_row_may_move():
    assert parse_mdl(make_mdl(header_row=7), NAME).header_row == 7


def test_invalid_mdl_gives_the_specified_message():          # TEST 12
    bad = make_mdl(columns=["S.No", "DISCIPLINE", "DOCUMENT REF", "NOTES", "Doc Category", "Remarks"])
    with pytest.raises(MdlFormatError) as exc:
        parse_mdl(bad, NAME)
    assert str(exc.value) == "Invalid MDL format.\nRequired columns DOC NO and DESCRIPTION were not found."
    with pytest.raises(MdlFormatError):
        parse_mdl(b"not an excel file", NAME)
    with pytest.raises(MdlFormatError):
        parse_mdl(make_mdl(), "mdl.xls")


def test_revision_cover_sheet_then_filename_then_not_identified():
    assert (parse_mdl(make_mdl(revision=5), NAME).revision, parse_mdl(make_mdl(revision=5), NAME).revision_source) == ("R5", "cover sheet")
    r = parse_mdl(make_mdl(add_cover=False), "PC286_R7_MDL.xlsx")
    assert (r.revision, r.revision_source) == ("R7", "file name")
    r = parse_mdl(make_mdl(add_cover=False), "project_mdl.xlsx")
    assert r.revision is None                                 # never invented
    r = parse_mdl(make_mdl(revision=5), NAME)                 # cover says 5, file name says R4
    assert r.revision == "R5" and any("R4" in w for w in r.warnings)


def test_split_document_number():
    assert split_document_number("PC286-10P195-I0003-100") == {
        "job_no": "PC286", "project_no": "10P195", "doc_code": "I0003", "serial_no": "100"}
    assert split_document_number("ABC-123") is None and split_document_number("A--B-C") is None


# ── 10: duplicate DOC NO ────────────────────────────────────────────────────
def test_duplicate_doc_no_keeps_both_records_with_distinct_ids():
    r = parse_mdl(make_mdl(), NAME)
    dups = [x for x in r.records if x.document_number == "PC286-10P195-D0032-301"]
    assert len(dups) == 2 and dups[0].mdl_record_id != dups[1].mdl_record_id
    assert r.duplicate_doc_numbers == {"PC286-10P195-D0032-301": 2}


def test_record_id_is_stable_across_revisions_and_independent_of_row():
    a = parse_mdl(make_mdl(), NAME).records
    shifted = make_mdl(rows=[(0, "X", "PC286-10P195-Z9999-001", "Inserted above", "A", "")] + ROWS, revision=5)
    b = parse_mdl(shifted, "PC286_R5_MDL.xlsx").records
    ida = {x.document_number + x.document_name: x.mdl_record_id for x in a}
    idb = {x.document_number + x.document_name: x.mdl_record_id for x in b}
    assert all(ida[k] == idb[k] for k in ida)                  # rows shifted by one; ids identical


# ── real MDL (2, 3, 4, 10, 11) ──────────────────────────────────────────────
@pytest.mark.skipif(not REAL_MDL.exists(), reason="sample MDL not available")
def test_real_mdl_1186_records_one_duplicate_and_r4():
    r = parse_mdl(REAL_MDL.read_bytes(), REAL_MDL.name)
    assert len(r.records) == 1186 and r.section_rows_ignored == 86
    assert r.revision == "R4" and r.source_sheet == "MDL"
    assert r.duplicate_doc_numbers == {"PC286-10P195-D0032-301": 2}
    assert len({x.mdl_record_id for x in r.records}) == 1186
    names = {x.document_name for x in r.records}
    assert "PROCESS" not in names and "Design Basis" not in names     # section headings are never documents


@pytest.mark.skipif(not REAL_MDL.exists(), reason="sample MDL not available")
def test_real_mdl_search_and_duplicates_via_api(client):
    assert upload(client, REAL_MDL.read_bytes(), name=REAL_MDL.name).status_code == 200
    assert len(search(client, "PC286-10P195-D0032-301", limit=50)["items"]) == 2
    both = search(client, "fabrication drawings UF Feed")["items"]
    assert any(i["document_name"] == "Fabrication drawings of UF Feed Collection Tank" and i["duplicate_document_number"] for i in both)
    assert search(client, "instrument datasheet", discipline="Instrumentation")["total"] > 10
    everything = client.get(f"/api/projects/{PID}/coversheet/mdl/documents", params={"limit": 500}).json()
    assert everything["total"] == 1186


# ── API: upload / status / search / select / replace ────────────────────────
def test_upload_status_and_search(client):                      # TESTS 1, 3, 4
    r = upload(client)
    assert r.status_code == 200
    body = r.json()
    assert body["uploaded"] and body["valid"] and body["sheet"] == "MDL"
    assert body["document_count"] == 7 and body["revision"] == "R4" and body["file_name"] == NAME
    assert body["previous"] is None and body["replaced"] is False
    by_no = search(client, "I0003-100")
    assert [i["document_name"] for i in by_no["items"]] == ["Instrument Datasheets for New WWTP"]
    by_name = search(client, "instrument datasheets zld")
    assert [i["document_number"] for i in by_name["items"]] == ["PC286-10P195-I0003-300"]
    assert search(client, "mechanical")["total"] == 2           # by discipline
    assert search(client, "")["total"] == 7 and search(client, "no-such-thing")["total"] == 0
    assert search(client, "100%_")["total"] == 0                # LIKE wildcards are escaped


def test_select_document_returns_coversheet_mapping(client):    # TEST 5
    upload(client)
    rec = search(client, "I0003-100")["items"][0]
    got = client.get(f"/api/projects/{PID}/coversheet/mdl/document/{rec['mdl_record_id']}").json()
    assert got["document_number"] == "PC286-10P195-I0003-100"
    assert got["coversheet"]["fields"] == {
        "doc_title_short": "Instrument Datasheets for New WWTP",
        "doc_title_full": "Instrument Datasheets for New WWTP",
        "job_no": "PC286", "project_no": "10P195", "doc_code": "I0003", "serial_no": "100"}
    assert client.get(f"/api/projects/{PID}/coversheet/mdl/document/mdl-nope").status_code == 404


def test_duplicate_selection_warns_and_does_not_pick_first(client):     # TEST 10
    upload(client)
    items = search(client, "D0032-301")["items"]
    assert len(items) == 2 and all(i["duplicate_document_number"] for i in items)
    chosen = next(i for i in items if "RO Feed" in i["document_name"])
    got = client.get(f"/api/projects/{PID}/coversheet/mdl/document/{chosen['mdl_record_id']}").json()
    assert got["document_name"] == "Fabrication drawings of RO Feed Tank"
    assert any("Multiple MDL entries" in w for w in got["coversheet"]["warnings"])


def test_invalid_upload_is_rejected_and_keeps_previous_mdl(client):
    upload(client)
    bad = make_mdl(columns=["S.No", "DISCIPLINE", "REF", "NOTES", "Doc Category", "Remarks"])
    r = upload(client, bad)
    assert r.status_code == 422 and "Required columns DOC NO and DESCRIPTION were not found." in r.json()["detail"]
    assert client.get(f"/api/projects/{PID}/coversheet/mdl").json()["document_count"] == 7


def test_reupload_replaces_and_revalidates_selection(client):   # TESTS 7, 8, 9
    upload(client)
    keep = search(client, "I0003-100")["items"][0]["mdl_record_id"]
    gone = search(client, "ZLD")["items"][0]["mdl_record_id"]

    rows_r5 = [r for r in ROWS if r[2] != "PC286-10P195-I0003-300"] + [
        (11, "Instrumentation", "PC286-10P195-I0003-400", "Instrument Datasheets for Pump House", "A", "")]
    kept = upload(client, make_mdl(rows=rows_r5, revision=5), name="PC286_R5_MDL.xlsx", selected=keep).json()
    assert kept["replaced"] and kept["previous"]["revision"] == "R4" and kept["revision"] == "R5"
    assert kept["selection"]["status"] == "preserved"
    assert kept["selection"]["record"]["document_name"] == "Instrument Datasheets for New WWTP"
    assert search(client, "pump house")["total"] == 1 and search(client, "ZLD")["total"] == 0   # list refreshed

    missing = upload(client, make_mdl(rows=rows_r5, revision=6), name="PC286_R6_MDL.xlsx", selected=gone).json()
    assert missing["selection"]["status"] == "missing"
    assert "no longer" not in missing["selection"]["message"] and "not available" in missing["selection"]["message"]
    assert missing["previous"]["revision"] == "R5"


def test_same_file_again_is_not_a_new_revision(client):
    upload(client)
    again = upload(client).json()
    assert again["unchanged"] is True and again["replaced"] is False


def test_mdl_is_project_specific(client):                       # section 14
    upload(client, make_mdl(), project=PID)
    other_rows = [(1, "Process", "PC999-10P204-P0001-001", "Only in the other project", "A", "")]
    upload(client, make_mdl(rows=other_rows), project=OTHER)
    assert search(client, "Only in the other", project=PID)["total"] == 0
    assert search(client, "Only in the other", project=OTHER)["total"] == 1
    rec_other = search(client, "", project=OTHER)["items"][0]["mdl_record_id"]
    assert client.get(f"/api/projects/{PID}/coversheet/mdl/document/{rec_other}").status_code == 404
    assert client.get("/api/projects/10P777/coversheet/mdl").json() == {"uploaded": False, "project_id": "10P777"}


def test_remove_mdl_restores_manual_mode(client):
    upload(client)
    assert client.delete(f"/api/projects/{PID}/coversheet/mdl").json()["removed"] is True
    assert client.get(f"/api/projects/{PID}/coversheet/mdl").json()["uploaded"] is False
    assert search(client, "")["total"] == 0


def test_bad_project_id_rejected(client):
    assert client.get("/api/projects/..%2Fx/coversheet/mdl").status_code in (400, 404)
    assert upload(client, project="bad id!").status_code in (400, 404)


# ── 6 / 13: generated file contains the MDL values ──────────────────────────
import json  # noqa: E402

REVISIONS = [{"rev_no": "0", "date": "26.10.2026", "description": "ISSUED FOR APPROVAL",
              "prepared_by": "AB", "checked_by": "CD", "approved_by": "EF"}]


def generate(client, project=None, record=None, info=None, overrides=None):
    project_info = {"client": "GAIL", "job_no": "MANUAL", "project_no": "MANUAL", "doc_code": "MANUAL",
                    "serial_no": "000", "unit_no": "200", "doc_title_short": "MANUAL TITLE",
                    "doc_title_full": "MANUAL FULL TITLE", "discipline": "INSTRUMENTATION", **(info or {})}
    payload = {"project_info": project_info, "revisions": REVISIONS}
    if overrides:
        payload["cell_overrides"] = overrides
    form = {"data": json.dumps(payload)}
    if project:
        form["project_id"] = project
    if record:
        form["mdl_record_id"] = record
    r = client.post("/api/coversheet/generate", data=form)
    return r


def download_workbook(client):
    r = client.get("/api/coversheet/download")
    assert r.status_code == 200
    return openpyxl.load_workbook(io.BytesIO(r.content)), r


def test_generated_file_contains_selected_mdl_values(client):    # TEST 6, acceptance #22
    upload(client)
    rec = search(client, "I0003-100")["items"][0]
    r = generate(client, PID, rec["mdl_record_id"])
    assert r.status_code == 200, r.text
    assert r.json()["mdl"]["document_number"] == "PC286-10P195-I0003-100"
    wb, resp = download_workbook(client)
    ws = wb.active
    assert ws["G2"].value == "Instrument Datasheets for New WWTP"                 # title block
    assert "Instrument Datasheets for New WWTP" in ws["A10"].value                # DOCUMENT TITLE box
    assert (ws["G4"].value, ws["J4"].value, ws["K4"].value, ws["M4"].value) == ("PC286", "10P195", "I0003", "100")
    assert "-".join([ws["G4"].value, ws["J4"].value, ws["K4"].value, ws["M4"].value]) == "PC286-10P195-I0003-100"
    assert ws["H4"].value == "200"                                                # unit no. untouched
    assert wb.properties.title == "PC286-10P195-I0003-100 - Instrument Datasheets for New WWTP"
    assert "PC286-10P195-I0003-100_CoverSheet.xlsx" in resp.headers["content-disposition"]
    assert ws["F6"].value == "GAIL" and ws["N4"].value == "0"                     # rest of the sheet intact


def test_mdl_overrides_manual_values_and_canvas_text_edits(client):
    upload(client)
    rec = search(client, "ZLD")["items"][0]
    edits = {"G2": {"value": "SOMEONE TYPED THIS"}, "K4": {"value": "HACKED", "alignment": {"horizontal": "center"}}}
    assert generate(client, PID, rec["mdl_record_id"], overrides=edits).status_code == 200
    ws = download_workbook(client)[0].active
    assert ws["G2"].value == "Instrument Datasheets for ZLD Unit" and ws["K4"].value == "I0003"
    assert ws["K4"].alignment.horizontal == "center"                              # formatting override kept


def test_duplicate_document_number_generates_the_selected_description(client):  # TEST 10
    upload(client)
    chosen = next(i for i in search(client, "D0032-301")["items"] if "RO Feed" in i["document_name"])
    r = generate(client, PID, chosen["mdl_record_id"])
    assert r.status_code == 200 and r.json()["mdl"]["warnings"]
    assert download_workbook(client)[0].active["G2"].value == "Fabrication drawings of RO Feed Tank"


def test_no_mdl_keeps_the_manual_workflow(client):               # TEST 13
    assert generate(client).status_code == 200                    # no project at all
    ws = download_workbook(client)[0].active
    assert (ws["G2"].value, ws["K4"].value, ws["M4"].value) == ("MANUAL TITLE", "MANUAL", "000")
    r = generate(client, project="10P555")                        # project without an MDL
    assert r.status_code == 200 and r.json()["mdl"] is None


def test_generation_blocked_until_a_document_is_selected(client):    # section 21
    upload(client)
    r = generate(client, PID)
    assert r.status_code == 400 and "Select the document from the MDL" in r.json()["detail"]
    r = generate(client, PID, "mdl-doesnotexist")
    assert r.status_code == 409 and "Select a document again" in r.json()["detail"]


def test_cannot_use_another_projects_record(client):
    upload(client, make_mdl(), project=PID)
    upload(client, make_mdl(rows=[(1, "P", "PC1-1-X0001-001", "Other project doc", "A", "")]), project=OTHER)
    foreign = search(client, "", project=OTHER)["items"][0]["mdl_record_id"]
    assert generate(client, PID, foreign).status_code == 409


def test_generation_is_audited_with_mdl_source(client, db):
    from app.services import document_registry as reg
    upload(client)
    rec = search(client, "I0003-100")["items"][0]
    generate(client, PID, rec["mdl_record_id"])
    actions = {a["action"]: a for a in reg.audit_trail(db, PID)}
    assert "mdl_uploaded" in actions and actions["coversheet_generated_from_mdl"]["details"]["mdl_file"] == NAME


def test_search_is_plural_tolerant(client):
    rows = [(1, "Instrumentation", "PC286-10P195-I0003-001", "Instrument Datasheet for PLC", "A", ""),
            (2, "Instrumentation", "PC286-10P195-I0003-002", "Instrument Datasheets for RTU", "A", "")]
    upload(client, make_mdl(rows=rows))
    assert search(client, "datasheets")["total"] == 2           # plural finds the singular spelling too
    assert search(client, "datasheet")["total"] == 2
    assert search(client, "instrument datasheets plc")["total"] == 1

"""
mdl_parser.py - parse a project's Master Document List (MDL) workbook.

The MDL is the authoritative source of document numbers and titles for the
Coversheet Generator:

    MDL "DOC NO"       -> Coversheet document number
    MDL "DESCRIPTION"  -> Coversheet document title

Rules (see the feature specification):
  * Data comes from the "MDL" sheet. The "Cover" sheet is never used for
    document records (it is read only to find a formal MDL revision).
  * Columns are found by header NAME, not by position.
  * A row is a document record only when DOC NO *and* DESCRIPTION are both
    populated. Section/category heading rows are ignored.
  * S.No is informational only (it may hold formula errors / odd sequences).
  * DOC NO is NOT unique. Identity is DOC NO + DESCRIPTION, never the Excel row.
"""
from __future__ import annotations

import hashlib
import io
import re
from dataclasses import dataclass, field
from typing import Any, Optional

INVALID_MDL_MESSAGE = ("Invalid MDL format.\n"
                       "Required columns DOC NO and DESCRIPTION were not found.")

_HEADER_SCAN_ROWS = 25
_ERROR_VALUES = {"#REF!", "#N/A", "#VALUE!", "#NAME?", "#DIV/0!", "#NULL!", "#NUM!"}

# Normalised header text -> logical field
_HEADER_ALIASES = {
    "doc_no": {"DOCNO", "DOCNUMBER", "DOCUMENTNO", "DOCUMENTNUMBER", "DOCNO.", "DOCUMENTNUMBERS"},
    "description": {"DESCRIPTION", "DOCDESCRIPTION", "DOCUMENTDESCRIPTION", "DOCUMENTTITLE",
                    "DOCUMENTNAME", "DOCTITLE", "TITLE"},
    "discipline": {"DISCIPLINE"},
    "category": {"DOCCATEGORY", "DOCUMENTCATEGORY", "CATEGORY"},
    "remarks": {"REMARKS", "REMARK"},
    "s_no": {"SNO", "SLNO", "SRNO", "SERIALNO", "SERIALNUMBER"},
}


class MdlFormatError(Exception):
    """The workbook is not a usable MDL (message is shown to the user as-is)."""


@dataclass
class MdlRecord:
    mdl_record_id: str
    s_no: str
    discipline: str
    document_number: str
    document_name: str
    document_category: str
    remarks: str
    source_file: str
    source_sheet: str


@dataclass
class MdlParseResult:
    source_file: str
    source_sheet: str
    header_row: int
    revision: Optional[str]
    revision_source: Optional[str]          # "cover sheet" | "file name" | None
    records: list[MdlRecord] = field(default_factory=list)
    section_rows_ignored: int = 0
    duplicate_doc_numbers: dict[str, int] = field(default_factory=dict)
    warnings: list[str] = field(default_factory=list)


# ── helpers ──────────────────────────────────────────────────────────────────
def _norm_header(value: Any) -> str:
    return re.sub(r"[^A-Z0-9]", "", str(value).upper()) if value not in (None, "") else ""


def _clean(value: Any) -> str:
    """Cell value -> single-line trimmed text ('' for blanks and Excel error values)."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = re.sub(r"\s+", " ", str(value).replace(" ", " ")).strip()
    return "" if text.upper() in _ERROR_VALUES else text


def record_id_for(doc_no: str, description: str, occurrence: int = 1) -> str:
    """Stable identity from DOC NO + DESCRIPTION (never the Excel row number)."""
    norm_desc = re.sub(r"\s+", " ", description).strip().lower()
    key = doc_no.strip().lower() + "|" + norm_desc
    rid = "mdl-" + hashlib.sha1(key.encode("utf-8")).hexdigest()[:16]
    return rid if occurrence == 1 else f"{rid}-{occurrence}"


def split_document_number(doc_no: str) -> Optional[dict[str, str]]:
    """
    WABAG document numbers look like  PC286-10P195-I0003-100  =
    JOB NO - PROJECT NO - DOC CODE - SERIAL NO.  Returns None if the number does
    not have exactly those four hyphen-separated parts (nothing is guessed).
    """
    parts = [p.strip() for p in str(doc_no).strip().split("-")]
    if len(parts) != 4 or not all(parts):
        return None
    return {"job_no": parts[0], "project_no": parts[1], "doc_code": parts[2], "serial_no": parts[3]}


def _find_header(rows: list[tuple]) -> Optional[tuple[int, dict[str, int]]]:
    """(1-based row number, {field: 0-based column}) of the first row holding DOC NO + DESCRIPTION."""
    for offset, row in enumerate(rows):
        cols: dict[str, int] = {}
        for idx, cell in enumerate(row):
            norm = _norm_header(cell)
            if not norm:
                continue
            for fld, aliases in _HEADER_ALIASES.items():
                if norm in aliases and fld not in cols:
                    cols[fld] = idx
        if "doc_no" in cols and "description" in cols:
            return offset + 1, cols
    return None


def _pick_sheets(wb) -> list:
    """The 'MDL' sheet first; otherwise other candidates - never the 'Cover' sheet."""
    sheets = [ws for ws in wb.worksheets if ws.title.strip().lower() != "cover"]
    exact = [ws for ws in sheets if ws.title.strip().lower() == "mdl"]
    named = [ws for ws in sheets if "mdl" in ws.title.lower() and ws not in exact]
    rest = [ws for ws in sheets if ws not in exact and ws not in named]
    return exact + named + rest


def _revision_from_cover(wb) -> Optional[str]:
    """Formal revision on the 'Cover' sheet: a cell reading REV/REVISION with a value to its right."""
    ws = next((s for s in wb.worksheets if s.title.strip().lower() == "cover"), None)
    if ws is None:
        return None
    for row in ws.iter_rows(min_row=1, max_row=80, max_col=30, values_only=True):
        for idx, cell in enumerate(row):
            if _norm_header(cell) in ("REV", "REVISION", "REVNO", "REVNUMBER"):
                for nxt in row[idx + 1: idx + 6]:
                    val = _clean(nxt)
                    if val:
                        return _format_revision(val)
    return None


def _format_revision(raw: str) -> Optional[str]:
    raw = raw.strip()
    m = re.fullmatch(r"(?:R(?:EV)?\.?\s*)?(\d{1,3}|[A-Z])", raw, flags=re.IGNORECASE)
    if not m:
        return None
    token = m.group(1).upper()
    return f"R{token}"


def _revision_from_filename(filename: str) -> Optional[str]:
    m = re.search(r"(?:^|[_\-\s.])R(?:EV)?[_\-\s]?(\d{1,3})(?=[_\-\s.]|$)", filename, flags=re.IGNORECASE)
    return f"R{m.group(1)}" if m else None


# ── main entry point ─────────────────────────────────────────────────────────
def parse_mdl(file_bytes: bytes, filename: str = "MDL.xlsx") -> MdlParseResult:
    """Parse an MDL workbook. Raises MdlFormatError with a user-facing message."""
    import openpyxl

    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext not in ("xlsx", "xlsm"):
        raise MdlFormatError("Unsupported MDL file type. Upload the MDL as an .xlsx or .xlsm workbook.")
    try:
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    except Exception as exc:
        raise MdlFormatError(f"The MDL file could not be opened: {exc}")

    chosen = None
    for ws in _pick_sheets(wb):
        head_rows = list(ws.iter_rows(min_row=1, max_row=_HEADER_SCAN_ROWS, max_col=40, values_only=True))
        found = _find_header(head_rows)
        if found:
            chosen = (ws, *found)
            break
    if chosen is None:
        wb.close()
        raise MdlFormatError(INVALID_MDL_MESSAGE)

    ws, header_row, cols = chosen
    max_col = max(cols.values()) + 1

    revision = _revision_from_cover(wb)
    revision_source = "cover sheet" if revision else None
    file_rev = _revision_from_filename(filename)
    warnings: list[str] = []
    if revision and file_rev and revision != file_rev:
        warnings.append(f"File name suggests {file_rev} but the cover sheet states {revision}; {revision} is used.")
    if not revision and file_rev:
        revision, revision_source = file_rev, "file name"

    def cell(row: tuple, fld: str) -> str:
        idx = cols.get(fld)
        return _clean(row[idx]) if idx is not None and idx < len(row) else ""

    records: list[MdlRecord] = []
    seen: dict[tuple[str, str], int] = {}
    doc_no_count: dict[str, int] = {}
    sections = 0
    for row in ws.iter_rows(min_row=header_row + 1, max_col=max_col, values_only=True):
        doc_no, desc = cell(row, "doc_no"), cell(row, "description")
        if not (doc_no and desc):
            if any(_clean(v) for v in row):      # a heading row such as "PROCESS" / "Design Basis"
                sections += 1
            continue
        key = (doc_no.lower(), desc.lower())
        seen[key] = seen.get(key, 0) + 1
        doc_no_count[doc_no] = doc_no_count.get(doc_no, 0) + 1
        records.append(MdlRecord(
            mdl_record_id=record_id_for(doc_no, desc, seen[key]),
            s_no=cell(row, "s_no"),
            discipline=cell(row, "discipline"),
            document_number=doc_no,
            document_name=desc,
            document_category=cell(row, "category"),
            remarks=cell(row, "remarks"),
            source_file=filename,
            source_sheet=ws.title,
        ))
    wb.close()

    if not records:
        raise MdlFormatError("The MDL sheet has the expected columns but no document records "
                             "(rows with both DOC NO and DESCRIPTION).")

    duplicates = {d: n for d, n in doc_no_count.items() if n > 1}
    if duplicates:
        warnings.append(f"{len(duplicates)} document number(s) appear more than once; all entries were kept.")
    exact = sum(1 for n in seen.values() if n > 1)
    if exact:
        warnings.append(f"{exact} entr{'y' if exact == 1 else 'ies'} repeat both DOC NO and DESCRIPTION exactly.")

    return MdlParseResult(
        source_file=filename, source_sheet=ws.title, header_row=header_row,
        revision=revision, revision_source=revision_source, records=records,
        section_rows_ignored=sections, duplicate_doc_numbers=duplicates, warnings=warnings,
    )

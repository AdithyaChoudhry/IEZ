"""Process Data Requisition: type library, tag matching, prefill, Excel output, API."""
import io
import sys
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))

from utils.process_data_library import INSTRUMENT_TYPES, IDENTIFICATION_COLUMNS  # noqa: E402
from utils.process_data_matcher import match_tags, build_prefill_values, extract_line_size  # noqa: E402
from utils.process_data_generator import generate_process_data_requisition, column_header  # noqa: E402

ID_HEADERS = ["S.NO", "TAG NO", "EQUIPMENT DESCRIPTION", "P&ID NO.", "SERVICE", "EQUIPMENT NO", "LINE", "LINE SIZE"]


def _df():
    rows = [
        # tag, desc, pid, type name, measuring, service, line, equip, density, tmin, tmax
        ("200-LIT-001", "LEVEL TX 1", "P001", "LEVEL INDICATING TRANSMITTER", "ULTRASONIC TYPE", "CLARIFIER", "600-SEW-MS-1", "TK-101", 1000, 5, 40),
        ("200-LIT-002", "LEVEL TX 2", "P002", "RADAR LEVEL TRANSMITTER", "-", "EQ TANK", "2\"-SW-20101-B24-NI", "TK-102", "-", "", ""),
        ("200-FIT-001", "FLOW TX", "P003", "FLOW INDICATING TRANSMITTER", "ELECTROMAGNETIC TYPE", "FEED", "300-RW-1", "P-1", "", "", ""),
        ("200-PG-001", "PRESSURE GAUGE", "P004", "PRESSURE GAUGE", "BOURDON TYPE", "PUMP DISCHARGE", "-", "P-1", "", "", ""),
        ("200-AIT-001", "DO", "P005", "ANALYZER INDICATING TRANSMITTER - DO", "-", "AERATION", "-", "AT-1", "", "", ""),
        ("200-AIT-002", "PH", "P006", "ANALYZER INDICATING TRANSMITTER - Ph", "-", "AERATION", "-", "AT-2", "", "", ""),
        ("200-AIT-003", "N", "P007", "ANALYZER INDICATING TRANSMITTER - NITROGEN", "-", "AERATION", "-", "AT-3", "", "", ""),
        ("200-P-001", "PUMP", "P008", "PUMP", "-", "LIFT STATION", "-", "P-001", "", "", ""),
    ]
    return pd.DataFrame(rows, columns=["TAG NO", "EQUIPMENT DESCRIPTION", "P&ID NO.", "INSTRUMENT TYPE NAME",
                                       "INSTRUMENT MEASURING  TYPE", "SERVICE", "LINE", "EQUIPMENT NO",
                                       "DENSITY (kg/m3)", "TEMPERATURE MIN (deg C)", "TEMPERATURE MAX (deg C)"])


def _tags(df, key):
    return [r["tag_no"] for r in match_tags(df, key)[0]]


def _generate(df, key, selected=None, prefill=False, subtype=None, custom=None):
    rows, _ = match_tags(df, key) if not custom else (custom["rows"], [])
    if selected is not None:
        rows = [r for r in rows if r["tag_no"] in selected]
    pre = build_prefill_values(df, rows, key, prefill) if not custom else [{} for _ in rows]
    data, name = generate_process_data_requisition(rows, pre, key, subtype, "10P211", "IODB.xlsx",
                                                   custom["fields"] if custom else None)
    return load_workbook(io.BytesIO(data)), name, rows


def test_library_has_all_requested_types_and_field_counts():
    labels = {c.label for c in INSTRUMENT_TYPES.values()}
    assert len(INSTRUMENT_TYPES) >= 30
    for must in ("Level Transmitter", "Flow Transmitter", "Temperature Transmitter", "Pressure Transmitter",
                 "Level Gauge", "Pressure Gauge", "Temperature Gauge", "pH Analyzer", "H2S Analyzer",
                 "Biogas Analyzer", "Custom Instrument"):
        assert must in labels, must
    assert len(INSTRUMENT_TYPES["level_transmitter"].fields) == 49
    assert [c.name for c in IDENTIFICATION_COLUMNS] == ID_HEADERS


def test_matching_is_safe_and_type_specific():
    df = _df()
    assert _tags(df, "level_transmitter") == ["200-LIT-001", "200-LIT-002"]
    assert _tags(df, "flow_transmitter") == ["200-FIT-001"]
    assert _tags(df, "pressure_gauge") == ["200-PG-001"]
    assert _tags(df, "dissolved_oxygen_analyzer") == ["200-AIT-001"]
    assert _tags(df, "ph_analyzer") == ["200-AIT-002"]           # only pH, not every AIT tag
    assert _tags(df, "nitrate_analyzer") == []                   # NITROGEN is never silently called nitrate
    assert _tags(df, "pressure_transmitter") == []               # a pressure gauge is not a transmitter


def test_line_size_extraction():
    assert extract_line_size("600-SEW-MS-1") == "600 mm"
    assert extract_line_size('2"-SW-20101-B24-NI') == '2"'
    assert extract_line_size("-") == ""
    assert extract_line_size(None) == ""


def test_workbook_identification_columns_and_blank_process_cells():
    wb, name, rows = _generate(_df(), "level_transmitter")
    assert name == "Process_Data_Requisition_Level_Transmitter_10P211.xlsx"
    assert wb.sheetnames[:2] == ["PROCESS DATA REQUISITION", "INSTRUCTIONS"]
    ws = wb["PROCESS DATA REQUISITION"]
    assert ws["A1"].value == "PROCESS DATA REQUISITION"
    hdr = [c.value for c in ws[5]]
    assert hdr[:8] == ID_HEADERS
    cfg = INSTRUMENT_TYPES["level_transmitter"]
    assert hdr[8:] == [column_header(f) for f in cfg.fields]
    assert len(hdr) == 8 + 49
    data = [[c.value for c in r] for r in ws.iter_rows(min_row=6, max_row=ws.max_row)]
    assert [d[1] for d in data] == ["200-LIT-001", "200-LIT-002"]
    assert [d[0] for d in data] == [1, 2]
    assert data[0][2:8] == ["LEVEL TX 1", "P001", "CLARIFIER", "TK-101", "600-SEW-MS-1", "600 mm"]
    assert data[1][7] == '2"'
    for d in data:                               # nothing fabricated: every process cell blank
        assert all(v in (None, "") for v in d[8:])
    assert ws.freeze_panes == "A6" and ws.auto_filter.ref


def test_select_only_some_tags():
    wb, _, rows = _generate(_df(), "level_transmitter", selected={"200-LIT-002"})
    ws = wb["PROCESS DATA REQUISITION"]
    assert [r[1].value for r in ws.iter_rows(min_row=6, max_row=ws.max_row)] == ["200-LIT-002"]


def test_prefill_off_and_on():
    df = _df()
    rows, _ = match_tags(df, "level_transmitter")
    off = build_prefill_values(df, rows, "level_transmitter", False)
    assert all(v == "" for r in off for v in r.values())
    on = build_prefill_values(df, rows, "level_transmitter", True)
    assert on[0]["Fluid Density Normal"] == "1000"           # copied from IODB
    assert on[0]["Fluid Density Min"] == "" and on[0]["Fluid Density Max"] == ""   # single value is not spread to min/max
    assert on[0]["Process Temperature Min"] == "5" and on[0]["Process Temperature Max"] == "40"
    assert on[1]["Fluid Density Normal"] == ""                # IODB placeholder '-' is not copied
    assert on[0]["Minimum Level"] == ""                       # no IODB source -> stays blank


def test_flow_subtype_fields_only_when_selected():
    df = _df()
    plain, _, _ = _generate(df, "flow_transmitter")
    mag, _, _ = _generate(df, "flow_transmitter", subtype="Magnetic")
    h_plain = [c.value for c in plain["PROCESS DATA REQUISITION"][5]]
    h_mag = [c.value for c in mag["PROCESS DATA REQUISITION"][5]]
    assert not any("Minimum Conductivity Required" in str(h) for h in h_plain)
    assert any("Minimum Conductivity Required" in str(h) for h in h_mag)
    assert not any("Channel Shape" in str(h) for h in h_mag)


def test_each_type_generates_its_own_columns():
    seen = {}
    for key, cfg in INSTRUMENT_TYPES.items():
        if cfg.is_custom:
            continue
        names = [f.name for f in cfg.fields]
        assert names and len(names) == len(set(names)) or key in ("pressure_transmitter",), key
        seen[key] = names
    assert "Free Chlorine Minimum" in seen["chlorine_analyzer"]
    assert "Expected pH Minimum" in seen["ph_analyzer"]
    assert any(n.startswith("DO ") for n in seen["dissolved_oxygen_analyzer"])
    assert "Centre-to-Centre Distance" in seen["level_gauge"]
    assert "Pressure Pulsation" in seen["pressure_gauge"]
    assert not set(seen["level_gauge"]) >= set(seen["pressure_gauge"])


def test_custom_instrument_columns_are_user_defined():
    rows = [{"tag_no": "X-1", "equipment_description": "d", "pid_no": "p", "service": "s",
             "equipment_no": "e", "line": "l", "line_size": ""}]
    fields = [{"parameter": "Residual Hydrogen Peroxide", "unit": "mg/L"}]
    wb, _, _ = _generate(_df(), "custom_analyzer", custom={"rows": rows, "fields": fields})
    hdr = [c.value for c in wb["PROCESS DATA REQUISITION"][5]]
    assert hdr[8:] == ["Residual Hydrogen Peroxide Minimum (mg/L)", "Residual Hydrogen Peroxide Normal (mg/L)",
                       "Residual Hydrogen Peroxide Maximum (mg/L)", "Residual Hydrogen Peroxide Remarks"]


def test_real_1000_row_iodb_if_present():
    path = ROOT / "DJB_IODB.xlsx"
    if not path.exists():
        pytest.skip("DJB_IODB.xlsx not in repo")
    from utils.file_handler import read_iodb
    df, err = read_iodb(open(path, "rb"))
    assert err is None
    assert len(_tags(df, "level_transmitter")) == 96 + 12
    assert len(_tags(df, "flow_transmitter")) == 54
    assert len(_tags(df, "pressure_transmitter")) == 36
    assert len(_tags(df, "pressure_gauge")) == 42
    assert len(_tags(df, "ph_analyzer")) == 12
    assert len(_tags(df, "dissolved_oxygen_analyzer")) == 24
    assert len(_tags(df, "temperature_transmitter")) == 12


# ---------------------------------------------------------------- API ----
def _iodb_bytes(df):
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        df.to_excel(xw, sheet_name="IODB", index=False)
    return buf.getvalue()


@pytest.fixture()
def api():
    from types import SimpleNamespace
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from app.deps import get_current_user
    from app.routers.process_data import router
    app = FastAPI()
    app.include_router(router, prefix="/api")
    app.dependency_overrides[get_current_user] = lambda: SimpleNamespace(id=7, username="alice")
    return TestClient(app)


def _post_match(api, key, df=None, subtype=None):
    data = {"instrument_type": key, **({"subtype": subtype} if subtype else {})}
    return api.post("/api/process-data/match-tags", data=data,
                    files={"file": ("iodb.xlsx", _iodb_bytes(df if df is not None else _df()))})


def test_api_types_list(api):
    r = api.get("/api/process-data/instrument-types")
    assert r.status_code == 200
    keys = [t["key"] for t in r.json()["instrument_types"]]
    assert "level_transmitter" in keys and "custom_analyzer" in keys and len(keys) >= 30


def test_api_match_generate_download(api):
    m = _post_match(api, "level_transmitter")
    assert m.status_code == 200 and m.json()["tag_count"] == 2
    assert len(m.json()["process_columns"]) == 49
    g = api.post("/api/process-data/generate", data={
        "instrument_type": "level_transmitter", "selected_tags": '["200-LIT-001"]',
        "use_iodb_values": "true", "project_id": "10P211"},
        files={"iodb_file": ("iodb.xlsx", _iodb_bytes(_df()))})
    assert g.status_code == 200, g.text
    assert g.json()["filename"] == "Process_Data_Requisition_Level_Transmitter_10P211.xlsx"
    d = api.get("/api/process-data/download")
    assert d.status_code == 200
    ws = load_workbook(io.BytesIO(d.content))["PROCESS DATA REQUISITION"]
    assert [r[1].value for r in ws.iter_rows(min_row=6, max_row=ws.max_row)] == ["200-LIT-001"]


def test_api_errors_do_not_produce_empty_file(api):
    r = _post_match(api, "nitrate_analyzer")
    assert r.status_code == 400 and "No matching Nitrate Analyzer tags" in r.json()["detail"]
    r = api.post("/api/process-data/generate", data={"instrument_type": "level_transmitter", "selected_tags": "[]"},
                 files={"iodb_file": ("iodb.xlsx", _iodb_bytes(_df()))})
    assert r.status_code == 400
    r = api.post("/api/process-data/generate", data={"instrument_type": "level_transmitter", "selected_tags": '["NOPE"]'},
                 files={"iodb_file": ("iodb.xlsx", _iodb_bytes(_df()))})
    assert r.status_code == 400
    r = _post_match(api, "no_such_type")
    assert r.status_code == 400


def test_api_custom_type_pick_tags_and_parameters(api):
    m = _post_match(api, "custom_analyzer")
    assert m.status_code == 200 and m.json()["tag_count"] == 8
    g = api.post("/api/process-data/generate", data={
        "instrument_type": "custom_analyzer", "selected_tags": '["200-P-001"]', "custom_name": "H2O2 Analyzer",
        "custom_fields": '[{"parameter": "Residual Hydrogen Peroxide", "unit": "mg/L"}]'},
        files={"iodb_file": ("iodb.xlsx", _iodb_bytes(_df()))})
    assert g.status_code == 200, g.text
    ws = load_workbook(io.BytesIO(api.get("/api/process-data/download").content))["PROCESS DATA REQUISITION"]
    hdr = [c.value for c in ws[5]]
    assert hdr[8:] == ["Residual Hydrogen Peroxide Minimum (mg/L)", "Residual Hydrogen Peroxide Normal (mg/L)",
                       "Residual Hydrogen Peroxide Maximum (mg/L)", "Residual Hydrogen Peroxide Remarks"]
    assert "H2O2 Analyzer" in ws["A2"].value


def test_api_selected_process_columns_only(api):
    m = _post_match(api, "level_transmitter").json()
    assert len(m["process_fields"]) == 49 and m["process_fields"][1]["header"] == "Minimum Level (m)"
    pick = ["Minimum Level", "Remarks"]
    g = api.post("/api/process-data/generate", data={
        "instrument_type": "level_transmitter", "selected_tags": '["200-LIT-001"]',
        "selected_fields": __import__("json").dumps(pick)},
        files={"iodb_file": ("iodb.xlsx", _iodb_bytes(_df()))})
    assert g.status_code == 200, g.text
    ws = load_workbook(io.BytesIO(api.get("/api/process-data/download").content))["PROCESS DATA REQUISITION"]
    assert [c.value for c in ws[5]][8:] == ["Minimum Level (m)", "Remarks"]
    bad = api.post("/api/process-data/generate", data={
        "instrument_type": "level_transmitter", "selected_tags": '["200-LIT-001"]', "selected_fields": '["Nope"]'},
        files={"iodb_file": ("iodb.xlsx", _iodb_bytes(_df()))})
    assert bad.status_code == 400
    empty = api.post("/api/process-data/generate", data={
        "instrument_type": "level_transmitter", "selected_tags": '["200-LIT-001"]', "selected_fields": '[]'},
        files={"iodb_file": ("iodb.xlsx", _iodb_bytes(_df()))})
    assert empty.status_code == 400

"""
Loop Wiring (TEST) - independent experimental module.
Numbered comments map to the ACCEPTANCE TESTS in the feature specification.
"""
import io
import json
from pathlib import Path
from types import SimpleNamespace

import openpyxl
import pandas as pd
import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient

import app.models.loop_test  # noqa: F401  (registers the template table before the fixtures create it)
from app.deps import get_current_user, get_db
from app.routers.loop_wiring_test import router as lt_router
from app.services import loop_test_service as svc
from utils.loop_test import diagram, library, matcher
from utils.loop_test.excel_out import build_workbook, safe_sheet_name
from utils.loop_test.io_assignment import io_class, parse_io_assignment
from utils.loop_test.iodb_fields import parse_iodb, resolve_iodb_fields

ROOT = Path(__file__).resolve().parents[2]
REAL_IO = ROOT / "LOOP WIRING INPUT.XLSX"
REAL_IODB = Path("/Users/adithyachoudhrym/Downloads/iEZ Inputs/Cable Schedule/IODB.xlsx")

IODB_COLS = ["TAG NO", "LOOP NUMBER", "INSTRUMENT TYPE NAME", "SIGNAL DESCRIPTION", "SIGNAL I/O TYPE", "JUNCTION BOX TAG NO.",
             "INSTRUMENT TO JB CABLE TAG NO", "INSTRUMENT TO JB CABLE TYPE", "BRANCH CABLE LENGTH", "MAIN CABLE TAG NUMBER",
             "MAIN CABLE TYPE", "MAIN CABLE LENGTH(MTRS)", "IS/NIS TYPE", "HART", "POWER SUPPLY", "2 WIRE / 4 WIRE"]


def iodb_row(tag, loop, sig, jb="JB-1", itype="LEVEL TRANSMITTER", length=27, hart="Y"):
    return [tag, loop, itype, itype, sig, jb, f"IC-{tag}", "1PX1.5 Sqmm", length, f"IC-{jb}", "12PX1.5 Sqmm", 100, "NIS", hart, "24V DC", "2 WIRE"]


def make_iodb(rows) -> bytes:
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        pd.DataFrame(rows, columns=IODB_COLS).to_excel(w, sheet_name="IODB - Table 1", index=False)
    return buf.getvalue()


IO_HEAD = ["S.No", "Tag Number", "Loop number", "JB No.", "JB TB1", "JB TB2", "IO Type", "Rack / Node", "Slot", "Channel No",
           "Module No", "Terminal Strip number", "Terminal No(CMN)", "Terminal No(Signal)", "Wire Type"]


def io_row(tag, loop, io_type, jb="JB-1", rack="1", slot="2", ch="3", module="MOD-A", strip="X101"):
    return [1, tag, loop, jb, "1", "2", io_type, rack, slot, ch, module, strip, "4", "5", "2 WIRE"]


def make_io(rows, header_row=3) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "IO PANEL-01"
    ws["A1"] = "PROJECT IO ASSIGNMENT"
    for c, h in enumerate(IO_HEAD, 1):
        ws.cell(header_row, c, h)
    for r, row in enumerate(rows, header_row + 1):
        for c, v in enumerate(row, 1):
            ws.cell(r, c, v)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.fixture()
def client(engine, db, store):
    svc._CACHE.clear()
    app = FastAPI()
    app.include_router(lt_router, prefix="/api")
    app.dependency_overrides[get_current_user] = lambda: SimpleNamespace(id=1, username="alice", email="a@x")
    app.dependency_overrides[get_db] = lambda: db
    return TestClient(app)


def new_session(client, iodb_rows, io_rows, project="10P195"):
    r = client.post("/api/loop-wiring-test/sessions", data={"project_id": project, "iodb_revision": "R2"},
                    files={"iodb_file": ("IODB.xlsx", make_iodb(iodb_rows))})
    assert r.status_code == 201, r.text
    sid = r.json()["session_id"]
    r = client.post(f"/api/loop-wiring-test/sessions/{sid}/io", data={"io_revision": "A"},
                    files={"io_file": ("IO.xlsx", make_io(io_rows))})
    assert r.status_code == 200, r.text
    return sid, f"/api/loop-wiring-test/sessions/{sid}"


def design(client, base, tag, loop_type, params=None, signals=None):
    r = client.post(base + "/config/default", json={"tag": tag, "loop_type": loop_type, "params": params or {},
                                                    "signals": signals or []})
    assert r.status_code == 200, r.text
    return r.json()


def preview(client, base, cfg):
    r = client.post(base + "/preview", json={"config": cfg})
    assert r.status_code == 200, r.text
    return r.json()


def failed(check_list):
    return {c["id"]: c for c in check_list["checks"] if c["status"] != "PASS"}


# ── parsers ─────────────────────────────────────────────────────────────────
def test_iodb_fields_resolve_semantically_and_missing_ones_are_unavailable():
    r = resolve_iodb_fields(["TAG NO", "MAIN CABLE LENGTH (Mtrs)", "JB TERMINAL +", "JB TERMINAL -", "IS/NIS TYPE", "REDUNDANT I O"])
    assert r["MAIN_CABLE_LENGTH"] == "MAIN CABLE LENGTH (Mtrs)" and r["IS_NIS"] == "IS/NIS TYPE"
    assert r["JB_TERMINAL_PLUS"] == "JB TERMINAL +" and r["JB_TERMINAL_MINUS"] == "JB TERMINAL -"      # '+' and '-' stay distinct
    assert r["ROUTE_NUMBER"] is None and r["ZENER_BARRIER"] is None                                    # unavailable, not an error


def test_iodb_without_a_tag_column_is_rejected_with_a_clear_message():
    buf = io.BytesIO()
    pd.DataFrame({"X": [1]}).to_excel(buf, index=False)
    with pytest.raises(ValueError, match="IODB"):
        parse_iodb(buf.getvalue())


def test_io_assignment_header_row_is_detected_not_assumed():
    data = parse_io_assignment(make_io([io_row("T-1", "L-1", "AI101")], header_row=3))
    assert data.sheets[0]["header_row"] == 3 and data.mapping["TAG_NO"]["column"] == "B"
    assert data.mapping["RACK"]["header"] == "Rack / Node" and data.mapping["TERMINAL_STRIP"]["column"] == "L"
    assert parse_io_assignment(make_io([io_row("T-1", "L-1", "AI101")], header_row=9)).sheets[0]["header_row"] == 9


def test_io_assignment_mapping_can_be_corrected_by_the_engineer():
    wb = openpyxl.Workbook()
    ws = wb.active
    for c, h in enumerate(["Instrument", "Loop", "Cabinet", "Position", "Point"], 1):
        ws.cell(1, c, h)
    ws.append(["T-1", "L-1", "RACK-9", "7", "12"])
    buf = io.BytesIO(); wb.save(buf)
    with pytest.raises(ValueError):
        parse_io_assignment(buf.getvalue())                                  # nothing recognisable -> clear error
    ws["A1"] = "Tag Number"; ws["B1"] = "Loop number"; ws["C1"] = "Rack"
    buf = io.BytesIO(); wb.save(buf)
    data = parse_io_assignment(buf.getvalue(), mapping_override={"SLOT": "D", "CHANNEL": "E"})
    f = data.rows[0]["fields"]
    assert (f["RACK"], f["SLOT"], f["CHANNEL"]) == ("RACK-9", "7", "12")


def test_spare_rows_are_ignored_in_the_tag_count():
    data = parse_io_assignment(make_io([io_row("T-1", "L-1", "AI101"), io_row("SPARE", "SPARE", "AI101")]))
    assert data.summary()["tags"] == 1 and sum(r["is_spare"] for r in data.rows) == 1


def test_io_class_extraction():
    assert (io_class("AI101"), io_class("DI-03"), io_class("do2"), io_class("900A16"), io_class("")) == ("AI", "DI", "DO", "", "")


@pytest.mark.skipif(not (REAL_IO.exists() and REAL_IODB.exists()), reason="sample project files not available")
def test_real_sample_files_parse_and_match():
    io_data = parse_io_assignment(REAL_IO.read_bytes(), REAL_IO.name)
    iodb = parse_iodb(REAL_IODB.read_bytes(), "IODB.xlsx")
    assert io_data.summary()["tags"] == 12 and len(iodb.rows) == 419
    res = matcher.match_all(iodb.rows, io_data.rows)
    by = {e["tag"]: e for e in res["entries"] if e["in_io"]}
    assert by["200-LIT-20201"]["status"] == "MATCHED"
    assert by["200-PIT-20101"]["status"] == "CONFLICT"                       # IODB JB-003 vs I/O sheet JB-001 (real finding)
    assert by["200-FIT-21401"]["status"] == "REVIEW REQUIRED"                # look-alike tag: suggestion only


# ── matching ────────────────────────────────────────────────────────────────
def mk(iodb_rows, io_rows):
    iodb = parse_iodb(make_iodb(iodb_rows))
    io_ = parse_io_assignment(make_io(io_rows))
    return iodb, io_


def test_match_hierarchy_exact_then_normalised():
    iodb, io_ = mk([iodb_row("200-LIT-1", "L-1", "AI"), iodb_row("200-LIT-2", "L-2", "AI")],
                   [io_row("200-LIT-1", "L-1", "AI1"), io_row(" 200 - lit-2 ", "L-2", "AI2")])
    e = {x["tag"]: x for x in matcher.match_all(iodb.rows, io_.rows)["entries"]}
    assert e["200-LIT-1"]["method"] == "exact tag" and e["200-LIT-2"]["method"] == "normalised tag"
    assert e["200-LIT-2"]["status"] == "MATCHED"


def test_not_found_either_way():
    iodb, io_ = mk([iodb_row("T-1", "L-1", "AI"), iodb_row("PUMP-STATION-77", "L-9", "AI")],
                   [io_row("T-1", "L-1", "AI1"), io_row("ZZ-424242", "L-8", "AI2")])
    res = matcher.match_all(iodb.rows, io_.rows)
    st = {e["tag"]: e["status"] for e in res["entries"]}
    assert st["PUMP-STATION-77"] == "IO NOT FOUND" and st["ZZ-424242"] == "IODB NOT FOUND" and st["T-1"] == "MATCHED"
    assert res["summary"]["not_matched"] == 2


def test_loop_number_and_look_alikes_are_suggestions_never_automatic():
    iodb, io_ = mk([iodb_row("200-LIT-20201A", "L-20201", "AI")], [io_row("200-LIT-20201", "L-20201", "AI1")])
    first = matcher.match_all(iodb.rows, io_.rows)["entries"][0]
    assert first["status"] == "REVIEW REQUIRED" and first["suggestions"][0]["iodb_tag"] == "200-LIT-20201A"
    assert not first["in_iodb"]
    dec = {"suggestions": {"200-LIT-20201": {"iodb_tag": "200-LIT-20201A", "action": "accept"}}}
    acc = matcher.match_all(iodb.rows, io_.rows, dec)["entries"][0]
    assert acc["status"] == "MATCHED" and "engineer-approved" in acc["method"]
    rej = matcher.match_all(iodb.rows, io_.rows, {"suggestions": {"200-LIT-20201": {"iodb_tag": "200-LIT-20201A", "action": "reject"}}})
    assert [e["status"] for e in rej["entries"] if e["tag"] == "200-LIT-20201"] == ["IODB NOT FOUND"]


def test_signal_and_jb_conflicts_block_until_decided():           # 12
    iodb, io_ = mk([iodb_row("T-1", "L-1", "AI", jb="JB-1")], [io_row("T-1", "L-1", "DI101", jb="JB-2")])
    e = matcher.match_all(iodb.rows, io_.rows)["entries"][0]
    assert e["status"] == "CONFLICT" and {c["field"] for c in e["conflicts"]} == {"SIGNAL_TYPE", "JB_TAG"}
    sig = next(c for c in e["conflicts"] if c["field"] == "SIGNAL_TYPE")
    assert (sig["iodb"], sig["io"]) == ("AI", "DI")
    dec = {"conflicts": {"T-1": {"SIGNAL_TYPE": "io", "JB_TAG": "iodb"}}}
    assert matcher.match_all(iodb.rows, io_.rows, dec)["entries"][0]["status"] == "MATCHED"


def test_multi_signal_tags_are_not_duplicates_but_true_duplicates_are():
    rows_iodb = [iodb_row("MTR-1", "M-1", s) for s in ("DO", "DI")]
    rows_io = [io_row("MTR-1", "M-1", "DO1", ch="1"), io_row("MTR-1", "M-1", "DI1", ch="2")]
    e = matcher.match_all(*[x.rows for x in mk(rows_iodb, rows_io)])["entries"][0]
    assert e["status"] == "MATCHED" and e["signals"] == 2
    dup_io = [io_row("T-1", "L-1", "AI1", ch="5"), io_row("T-1", "L-1", "AI1", ch="5")]
    assert matcher.match_all(*[x.rows for x in mk([iodb_row("T-1", "L-1", "AI")], dup_io)])["entries"][0]["status"] == "DUPLICATE"


def test_merged_data_keeps_sources_apart_and_never_overwrites_silently():
    iodb, io_ = mk([iodb_row("T-1", "L-1", "AI", jb="JB-IODB")], [io_row("T-1", "L-1", "AI1", jb="JB-IO", rack="R7")])
    entry = matcher.match_all(iodb.rows, io_.rows)["entries"][0]
    d = matcher.merged_tag_data(entry, iodb.rows, io_.rows)
    assert d["fields"]["RACK"] == {**d["fields"]["RACK"], "value": "R7", "source": "I/O ASSIGNMENT"}
    assert d["fields"]["BRANCH_CABLE_TYPE"]["source"] == "IODB" and d["fields"]["JB_TAG"]["value"] == "JB-IODB"   # conflict: IODB kept until decided
    chosen = matcher.merged_tag_data(entry, iodb.rows, io_.rows, {"conflicts": {"T-1": {"JB_TAG": "io"}}})
    assert chosen["fields"]["JB_TAG"]["value"] == "JB-IO" and chosen["fields"]["JB_TAG"]["source"] == "I/O ASSIGNMENT"
    user = matcher.merged_tag_data(entry, iodb.rows, io_.rows, overrides={"PLC": "PLC-01"})
    assert user["fields"]["PLC"]["source"] == "USER CONFIGURATION"


# ── library ─────────────────────────────────────────────────────────────────
def test_rtd_wiring_choice_changes_the_conductors():                 # 5
    assert [len(library.new_config("RTD", {"rtd_wires": n})["links"][0]["wires"]) for n in ("2", "3", "4")] == [2, 3, 4]
    assert len(library.new_config("RTD", {})["links"][0]["wires"]) == 0          # not assumed


def test_motor_and_valve_signals_depend_on_type_and_are_never_preselected():
    assert library.new_config("MOTOR", {"starter": "VFD"})["signals"] == []
    vfd = {s["id"] for s in library.signal_catalog("MOTOR", {"starter": "VFD"})}
    dol = {s["id"] for s in library.signal_catalog("MOTOR", {"starter": "DOL"})}
    sd = {s["id"] for s in library.signal_catalog("MOTOR", {"starter": "STAR_DELTA"})}
    assert {"speed_reference", "speed_feedback"} <= vfd and not {"speed_reference"} & dol
    assert {"star_command", "delta_command"} <= sd and not {"star_command"} & vfd
    mov = {s["id"] for s in library.signal_catalog("VALVE", {"valve_type": "MOV"})}
    assert {"torque_trip", "open_command"} <= mov
    assert library.new_config("VALVE", {"valve_type": "MOV"})["signals"] == []


def test_planned_fieldbus_types_are_listed_but_not_creatable():
    assert library.LOOP_TYPES["COMM_MODBUS"]["status"] == "planned"
    with pytest.raises(ValueError):
        library.new_config("COMM_MODBUS")
    with pytest.raises(ValueError):
        library.new_config("NOPE")


def test_components_can_be_added_removed_and_links_follow():
    cfg = library.new_config("AI_2WIRE")
    cfg["components"].insert(2, library.new_component("zener_barrier", 9))
    library.reconcile_links(cfg)
    assert [c["kind"] for c in cfg["components"]] == ["field_instrument", "junction_box", "zener_barrier", "marshalling", "plc"]
    assert len(cfg["links"]) == 4 and all(l["wires"] for l in cfg["links"])


# ── API: the acceptance tests ───────────────────────────────────────────────
def base_data():
    iodb_rows = [iodb_row("LIT-1", "L-1", "AI"), iodb_row("XS-1", "X-1", "DI", itype="SWITCH"), iodb_row("SV-1", "S-1", "DO", itype="SOLENOID"),
                 iodb_row("FCV-1", "F-1", "AO", itype="CONTROL VALVE"), iodb_row("TE-1", "T-1", "RTD", itype="RTD")]
    io_rows = [io_row("LIT-1", "L-1", "AI101", ch="1"), io_row("XS-1", "X-1", "DI101", ch="2"), io_row("SV-1", "S-1", "DO101", ch="3"),
               io_row("FCV-1", "F-1", "AO101", ch="4"), io_row("TE-1", "T-1", "RTD1", ch="5")]
    return iodb_rows, io_rows


def test_01_to_04_standard_loops_match_and_generate(client):
    sid, base = new_session(client, *base_data())
    for tag, lt_id in (("LIT-1", "AI_2WIRE"), ("XS-1", "DI"), ("SV-1", "DO"), ("FCV-1", "AO")):
        cfg = design(client, base, tag, lt_id)
        pv = preview(client, base, cfg)
        assert pv["validation"]["can_generate"], (tag, failed(pv["validation"]))
        assert tag in pv["svg"] and "Rack: 1" in pv["svg"]
    g = client.post(base + "/generate", json={"tags": ["LIT-1"], "configs": {"LIT-1": design(client, base, "LIT-1", "AI_2WIRE")}}).json()
    assert g["generated"] == 1 and g["filename"] == "Loop_Wiring_Test_10P195_R2.xlsx"


def test_05_rtd_three_wire_appears_in_preview_and_workbook(client):
    sid, base = new_session(client, *base_data())
    cfg = design(client, base, "TE-1", "RTD", {"rtd_wires": "3"})
    pv = preview(client, base, cfg)
    assert "RTD 3 (lead compensation)" in pv["svg"] and pv["validation"]["can_generate"]
    bad = design(client, base, "TE-1", "RTD", {})
    v = client.post(base + "/validate", json={"config": bad}).json()
    assert not v["can_generate"] and "special" in failed(v)                      # wiring must be chosen


def motor_session(client, n_signals):
    names = ["DO", "DI", "DI", "AO", "AI"][:n_signals]
    iodb_rows = [iodb_row("P-101", "M-101", s, itype="MOTOR") for s in names]
    io_rows = [io_row("P-101", "M-101", f"{s}1", ch=str(i + 1)) for i, s in enumerate(names)]
    return new_session(client, iodb_rows, io_rows)


def test_06_motor_dol_is_one_loop_with_the_configured_signals(client):
    sid, base = motor_session(client, 3)
    cfg = design(client, base, "P-101", "MOTOR", {"starter": "DOL"}, ["start", "run_feedback", "fault"])
    assert [s["id"] for s in cfg["signals"]] == ["start", "run_feedback", "fault"]
    unlinked = client.post(base + "/validate", json={"config": cfg}).json()
    assert not unlinked["can_generate"] and "signal_link" in failed(unlinked)    # channels are never allocated automatically
    for i, s in enumerate(cfg["signals"]):
        s["io_index"] = i
    pv = preview(client, base, cfg)
    assert pv["validation"]["can_generate"], failed(pv["validation"])
    assert all(n in pv["svg"] for n in ("Start command", "Run feedback", "Fault"))
    g = client.post(base + "/generate", json={"tags": ["P-101"], "configs": {"P-101": cfg}}).json()
    assert g["generated"] == 1


def test_07_motor_vfd_combined_signals(client):
    sid, base = motor_session(client, 5)
    cfg = design(client, base, "P-101", "MOTOR", {"starter": "VFD"},
                 ["start", "run_feedback", "fault", "speed_reference", "speed_feedback"])
    for i, s in enumerate(cfg["signals"]):
        s["io_index"] = i
    pv = preview(client, base, cfg)
    assert pv["validation"]["can_generate"] and len(pv["model"]["signals"]) == 5
    assert "Speed reference" in pv["svg"]
    wrong = design(client, base, "P-101", "MOTOR", {"starter": "DOL"}, ["speed_reference"])       # not valid for DOL: ignored, none selected
    assert wrong["signals"] == []


def test_08_shielding_choices_are_represented(client):
    sid, base = new_session(client, *base_data())
    cfg = design(client, base, "LIT-1", "AI_2WIRE")
    for l in cfg["links"]:
        l["cable"]["shield_type"], l["cable"]["shield_termination"] = "individual", "jb"
    assert "Individual shield · JB end" in preview(client, base, cfg)["svg"]
    for l in cfg["links"]:
        l["cable"]["shield_type"], l["cable"]["shield_termination"] = "overall", "control_room"
    assert "Overall / common shield · Control-room end" in preview(client, base, cfg)["svg"]
    assert "Shield: not defined" in preview(client, base, design(client, base, "LIT-1", "AI_2WIRE"))["svg"]   # never silently chosen


def test_09_10_barrier_and_connector_appear_in_the_users_order(client):
    sid, base = new_session(client, *base_data())
    cfg = design(client, base, "LIT-1", "AI_2WIRE")
    cfg["components"].insert(2, library.new_component("zener_barrier", 20, "Zener barrier"))
    cfg["components"].insert(1, library.new_component("connector", 21, "Connector"))
    pv = preview(client, base, cfg)
    kinds = [c["kind"] for c in pv["config"]["components"]]
    assert kinds == ["field_instrument", "connector", "junction_box", "zener_barrier", "marshalling", "plc"]
    svg = pv["svg"]
    assert svg.index("CONNECTOR") < svg.index("JUNCTION BOX") < svg.index("ZENER BARRIER") < svg.index("MARSHALLING PANEL")


def test_11_missing_plc_data_blocks_and_nothing_is_drawn_for_it(client):
    iodb_rows = [iodb_row("LIT-1", "L-1", "AI")]
    io_rows = [io_row("LIT-1", "L-1", "AI101", rack="", slot="", ch="")]
    sid, base = new_session(client, iodb_rows, io_rows)
    cfg = design(client, base, "LIT-1", "AI_2WIRE")
    pv = preview(client, base, cfg)
    f = failed(pv["validation"])
    assert not pv["validation"]["can_generate"] and {"rack", "slot", "channel"} <= set(f)
    assert pv["svg"].count("MISSING") >= 3
    g = client.post(base + "/generate", json={"tags": ["LIT-1"], "configs": {"LIT-1": cfg}}).json()
    assert g["generated"] == 0 and g["results"][0]["status"] == "BLOCKED" and "Rack" in g["results"][0]["reason"]
    cfg["overrides"] = {"RACK": "R1", "SLOT": "S3", "CHANNEL": "CH5"}                  # engineer supplies it explicitly (source USER)
    assert preview(client, base, cfg)["validation"]["can_generate"]


def test_missing_jb_is_not_invented(client):
    sid, base = new_session(client, [iodb_row("LIT-1", "L-1", "AI", jb="-")], [io_row("LIT-1", "L-1", "AI101", jb="")])
    cfg = design(client, base, "LIT-1", "AI_2WIRE")
    v = preview(client, base, cfg)["validation"]
    assert "jb" in failed(v) and not v["can_generate"]
    cfg["components"] = [c for c in cfg["components"] if c["kind"] != "junction_box"]    # removing the JB from the loop is the engineer's call
    assert preview(client, base, cfg)["validation"]["can_generate"]


def test_12_conflict_is_displayed_then_decided_then_generated(client):
    sid, base = new_session(client, [iodb_row("LIT-1", "L-1", "AI")], [io_row("LIT-1", "L-1", "DI101")])
    assert client.get(base + "/matches").json()["items"][0]["status"] == "CONFLICT"
    cfg = design(client, base, "LIT-1", "AI_2WIRE")
    v = preview(client, base, cfg)["validation"]
    assert "signal" in failed(v) and "AI" in failed(v)["signal"]["detail"] and "DI" in failed(v)["signal"]["detail"]
    r = client.post(base + "/decisions", json={"kind": "conflict", "tag": "LIT-1", "field": "SIGNAL_TYPE", "choice": "iodb"})
    assert r.json()["summary"]["conflicts"] == 0
    v2 = preview(client, base, cfg)["validation"]                 # engineer accepted the IODB (AI): generation may continue, with a warning
    assert v2["can_generate"] and failed(v2)["signal"]["status"] == "WARNING"
    client.post(base + "/decisions", json={"kind": "conflict", "tag": "LIT-1", "field": "SIGNAL_TYPE", "choice": "io"})
    v3 = preview(client, base, cfg)["validation"]                 # accepted the I/O Assignment (DI) but the loop is AI -> must choose a DI loop
    assert not v3["can_generate"] and "DI" in failed(v3)["signal"]["detail"]
    di = design(client, base, "LIT-1", "DI")
    assert preview(client, base, di)["validation"]["can_generate"]
    assert client.post(base + "/decisions", json={"kind": "conflict", "tag": "LIT-1", "field": "SIGNAL_TYPE", "choice": "bogus"}).status_code == 400


def test_13_ten_loops_ten_sheets(client):
    iodb_rows = [iodb_row(f"LIT-{i}", f"L-{i}", "AI") for i in range(1, 11)]
    io_rows = [io_row(f"LIT-{i}", f"L-{i}", "AI101", ch=str(i)) for i in range(1, 11)]
    sid, base = new_session(client, iodb_rows, io_rows)
    cfgs = {f"LIT-{i}": design(client, base, f"LIT-{i}", "AI_2WIRE") for i in range(1, 11)}
    g = client.post(base + "/generate", json={"tags": list(cfgs), "configs": cfgs}).json()
    assert g["generated"] == 10
    wb = openpyxl.load_workbook(io.BytesIO(client.get(base + "/download").content))
    assert wb.sheetnames == [f"LIT-{i}" for i in range(1, 11)]
    assert wb["LIT-3"]["A1"].value == "LOOP WIRING (TEST) — LIT-3"
    text = " ".join(str(c.value) for row in wb["LIT-3"].iter_rows() for c in row if c.value)
    assert "Channel: 3" in text and "Channel: 4" not in text                # each loop carries its own channel


def test_14_templates_save_reload_delete(client):
    sid, base = new_session(client, *base_data())
    cfg = design(client, base, "LIT-1", "AI_2WIRE", {"hart": True})
    cfg["components"].insert(2, library.new_component("zener_barrier", 30))
    for l in design(client, base, "LIT-1", "AI_2WIRE")["links"]:
        pass
    r = client.post("/api/loop-wiring-test/templates", json={"name": "AI 2-wire + barrier", "description": "IS AI loop", "config": cfg})
    assert r.status_code == 201
    lib = client.get("/api/loop-wiring-test/library").json()
    saved = next(t for t in lib["saved_templates"] if t["name"] == "AI 2-wire + barrier")
    assert saved["config"]["tag_no"] == "" and any(c["kind"] == "zener_barrier" for c in saved["config"]["components"])
    applied = client.post(base + "/config/from-template", json={"template_id": saved["id"], "tag": "XS-1"}).json()
    assert applied["tag_no"] == "XS-1" and any(c["kind"] == "zener_barrier" for c in applied["components"])
    builtin = client.post(base + "/config/from-template", json={"template_id": "RTD_4WIRE", "tag": "TE-1"}).json()
    assert len(builtin["links"][0]["wires"]) == 4
    assert client.delete(f"/api/loop-wiring-test/templates/{saved['db_id']}").status_code == 200
    assert not client.get("/api/loop-wiring-test/templates").json()


def test_15_production_loop_wiring_is_untouched():
    router_src = (ROOT / "backend/app/routers/loop_wiring.py").read_text()
    assert "loop_test" not in router_src and "loop_wiring_test" not in router_src
    gen_src = (ROOT / "utils/generators.py").read_text()
    assert "utils.loop_test" not in gen_src and "loop_test" not in gen_src.split("def generate_loop_wiring")[1].split("\ndef ")[0]
    from app.routers import loop_wiring
    assert loop_wiring.router.prefix == "/loop-wiring"


# ── safety: sessions, persistence, escaping ─────────────────────────────────
def test_sessions_survive_a_restart_and_are_private(client, db):
    sid, base = new_session(client, *base_data())
    svc._CACHE.clear()                                            # what a deploy / restart does to memory
    assert client.get(base).json()["io"]["file_name"] == "IO.xlsx"
    assert client.post(base + "/match").json()["summary"]["matched"] == 5
    app2 = FastAPI()
    app2.include_router(lt_router, prefix="/api")
    app2.dependency_overrides[get_current_user] = lambda: SimpleNamespace(id=99, username="mallory", email="m@x")
    app2.dependency_overrides[get_db] = lambda: db
    assert TestClient(app2).get(base).status_code == 400          # another user cannot read the session
    assert client.get("/api/loop-wiring-test/sessions/not-a-session").status_code == 400


def test_hostile_text_is_escaped_in_svg():
    cfg = library.new_config("AI_2WIRE", tag_no="T-1")
    cfg["components"][0]["tag"] = '<script>alert(1)</script>"&'
    data = {"fields": {"TAG_NO": {"value": "T-1", "source": "IODB", "label": "Tag", "available": True}}, "signals": [], "conflicts": []}
    svg = diagram.render_svg(diagram.build_diagram(cfg, data, {}))
    assert "<script>" not in svg and "&lt;script&gt;" in svg


def test_sheet_names_are_safe_and_unique():
    used: set[str] = set()
    names = [safe_sheet_name(t, used) for t in ("A/B", "A/B", "X" * 40, "[T]:1")]
    assert len(set(n.lower() for n in names)) == 4 and all(len(n) <= 31 and not set(n) & set("[]:*?/\\") for n in names)


def test_horizontal_orientation_renders_in_svg_and_excel():
    cfg = library.new_config("AI_2WIRE", tag_no="T-1")
    cfg["layout"]["orientation"] = "horizontal"
    data = {"fields": {"TAG_NO": {"value": "T-1", "source": "IODB", "label": "Tag", "available": True}}, "signals": [], "conflicts": []}
    model = diagram.build_diagram(cfg, data, {"project": "P"})
    assert "<svg" in diagram.render_svg(model)
    wb = openpyxl.load_workbook(io.BytesIO(build_workbook([{"tag": "T-1", "model": model}])))
    assert wb.sheetnames == ["T-1"]


def test_clone_binds_the_structure_to_each_tag_without_guessing_channels(client):
    iodb_rows = [iodb_row(f"LIT-{i}", f"L-{i}", "AI") for i in (1, 2)] + [iodb_row("P-1", "M-1", s, itype="MOTOR") for s in ("DO", "DI")]
    io_rows = [io_row("LIT-1", "L-1", "AI1", ch="1"), io_row("LIT-2", "L-2", "AI1", ch="9"),
               io_row("P-1", "M-1", "DO1", ch="1"), io_row("P-1", "M-1", "DI1", ch="2")]
    sid, base = new_session(client, iodb_rows, io_rows)
    cfg = design(client, base, "LIT-1", "AI_2WIRE")
    cfg["overrides"] = {"RACK": "R-ONLY-FOR-LIT-1"}
    c2 = client.post(base + "/config/clone", json={"config": cfg, "tag": "LIT-2"}).json()
    assert c2["tag_no"] == "LIT-2" and c2["loop_number"] == "L-2" and c2["overrides"] == {} and c2["signals"][0]["io_index"] == 0
    assert client.post(base + "/validate", json={"config": c2}).json()["can_generate"]
    motor = design(client, base, "P-1", "MOTOR", {"starter": "DOL"}, ["start", "stop", "run_feedback"])      # 3 signals, 2 I/O rows
    cm = client.post(base + "/config/clone", json={"config": motor, "tag": "P-1"}).json()
    assert [s["io_index"] for s in cm["signals"]] == [None, None, None]
    assert "signal_link" in failed(client.post(base + "/validate", json={"config": cm}).json())


def test_motor_segment_wires_follow_the_signals_until_the_engineer_edits_them(client):
    sid, base = motor_session(client, 3)
    cfg = design(client, base, "P-101", "MOTOR", {"starter": "DOL"}, ["start", "run_feedback", "fault"])
    assert all(len(l["wires"]) == 6 for l in cfg["links"])                       # one pair per signal
    assert cfg["links"][0]["wires"][0]["label"] == "Start command +"
    cfg["signals"] = cfg["signals"][:2]                                          # engineer drops a signal
    again = client.post(base + "/preview", json={"config": cfg}).json()["config"]
    assert all(len(l["wires"]) == 4 for l in again["links"])
    again["links"][0]["wires"] = again["links"][0]["wires"][:2]                   # engineer edits the first segment by hand
    again["links"][0]["auto_wires"] = False
    again["signals"] = again["signals"] + [{"id": "stop", "name": "Stop command", "signal_type": "DO",
                                            "direction": "system_to_field", "io_index": None}]
    final = client.post(base + "/preview", json={"config": again}).json()["config"]
    assert len(final["links"][0]["wires"]) == 2 and len(final["links"][1]["wires"]) == 6    # edited link keeps its wires

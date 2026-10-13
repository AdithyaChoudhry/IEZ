"""
Document registry service.

The registry is the single place that knows which controlled deliverables were
generated for a project, which source-file revision each came from, and which
revision the Coversheet is linked to.

Principles:
  * IODB (or the Loop Wiring input) is the engineering source.
  * Generators create deliverables and *register* them here - the generation
    logic itself stays in utils/ and is never duplicated.
  * The Coversheet only reads the registry. It never edits engineering data.
  * A new revision never silently replaces a linked one - the link is moved
    only when an engineer explicitly connects the new revision.
"""
from __future__ import annotations

import hashlib
import io
import json
import logging
import re
import sys
from datetime import datetime
from pathlib import Path
from typing import Any, Optional

from sqlalchemy import func
from sqlalchemy.orm import Session

from ..config import settings
from ..models.document_registry import (
    CoversheetLink, DocumentAuditLog, GeneratedDocument, Project, SourceVersion,
)

logger = logging.getLogger(__name__)

# ── Module catalogue ─────────────────────────────────────────────────────────
# `kind` is the source file the module is generated from; revisions of a
# document are compared against the latest revision of that same kind.
MODULES: dict[str, dict[str, Any]] = {
    "instrument_list": {"name": "Instrument List", "serial": 1, "kind": "iodb",       "count_label": "Records"},
    "io_list":         {"name": "I/O List",        "serial": 2, "kind": "iodb",       "count_label": "Signals"},
    "cable_schedule":  {"name": "Cable Schedule",  "serial": 3, "kind": "iodb",       "count_label": "Cables"},
    "loop_wiring":     {"name": "Loop Wiring",     "serial": 4, "kind": "loop_input", "count_label": "Loops"},
}
MODULE_ORDER = list(MODULES.keys())
SOURCE_KINDS = ("iodb", "loop_input")

# Extends (does not replace) any approval flow that already exists in the app.
STATUSES = ("DRAFT", "GENERATED", "UNDER_REVIEW", "APPROVED", "SUPERSEDED", "FAILED")
MANUAL_STATUSES = ("DRAFT", "GENERATED", "UNDER_REVIEW", "APPROVED")

_PROJECT_ID_RE = re.compile(r"^[A-Za-z0-9][A-Za-z0-9._-]{0,39}$")


class RegistryError(Exception):
    """Raised for expected, user-facing problems (bad project id, nothing to connect...)."""


# ── Helpers ──────────────────────────────────────────────────────────────────
def normalize_project_id(project_id: str) -> str:
    pid = (project_id or "").strip()
    if not _PROJECT_ID_RE.match(pid):
        raise RegistryError(
            "Invalid project id. Use 1-40 letters, digits, '.', '_' or '-' (e.g. 10P211)."
        )
    return pid


def _store_dir() -> Path:
    configured = getattr(settings, "DOCUMENT_STORE_DIR", "") or ""
    base = Path(configured) if configured else Path(__file__).resolve().parents[2] / "document_store"
    base.mkdir(parents=True, exist_ok=True)
    return base


def _safe_name(name: str) -> str:
    cleaned = re.sub(r"[^A-Za-z0-9._-]+", "_", name or "file").strip("._") or "file"
    return cleaned[:120]


def _sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _iso(dt: Optional[datetime]) -> Optional[str]:
    return dt.isoformat() + "Z" if dt else None


def _audit(db: Session, project_id: str, action: str, actor: str | None,
           module: str | None = None, document_id: int | None = None, **details: Any) -> None:
    db.add(DocumentAuditLog(
        project_id=project_id, module=module, document_id=document_id, action=action,
        actor=actor, details=json.dumps(details, default=str) if details else None,
    ))


def _ensure_utils_importable() -> None:
    """Make the shared `utils` package importable (same trick the routers use)."""
    here = Path(__file__).resolve()
    for p in (here.parents[2], here.parents[3], Path("/app")):
        if (p / "utils").is_dir() and str(p) not in sys.path:
            sys.path.insert(0, str(p))


# ── Projects & source versions ───────────────────────────────────────────────
def get_or_create_project(db: Session, project_id: str, name: str | None = None,
                          user: str | None = None) -> Project:
    pid = normalize_project_id(project_id)
    proj = db.query(Project).filter(Project.project_id == pid).first()
    if proj:
        if name and name.strip() and proj.name == proj.project_id:
            proj.name = name.strip()
        return proj
    proj = Project(project_id=pid, name=(name or pid).strip() or pid, created_by=user)
    db.add(proj)
    db.flush()
    return proj


def latest_source(db: Session, project_id: str, kind: str) -> Optional[SourceVersion]:
    return (db.query(SourceVersion)
            .filter(SourceVersion.project_id == project_id, SourceVersion.kind == kind)
            .order_by(SourceVersion.rev_no.desc()).first())


def ensure_source_version(db: Session, project_id: str, kind: str, filename: str,
                          data: bytes, user: str | None) -> SourceVersion:
    """Return the version row for this exact file, creating the next revision if it is new."""
    if kind not in SOURCE_KINDS:
        raise RegistryError(f"Unknown source kind '{kind}'.")
    digest = _sha256(data)
    existing = (db.query(SourceVersion)
                .filter(SourceVersion.project_id == project_id,
                        SourceVersion.kind == kind, SourceVersion.sha256 == digest).first())
    if existing:
        return existing

    last = db.query(func.max(SourceVersion.rev_no)).filter(
        SourceVersion.project_id == project_id, SourceVersion.kind == kind).scalar() or 0
    rev = last + 1
    folder = _store_dir() / "projects" / project_id / "sources"
    folder.mkdir(parents=True, exist_ok=True)
    path = folder / f"{kind}_rev{rev}_{digest[:12]}_{_safe_name(filename)}"
    path.write_bytes(data)

    sv = SourceVersion(project_id=project_id, kind=kind, rev_no=rev, sha256=digest,
                       filename=filename or "source.xlsx", file_path=str(path), created_by=user)
    db.add(sv)
    db.flush()
    _audit(db, project_id, "source_uploaded", user, kind=kind, rev=rev, filename=filename, sha256=digest)
    return sv


# ── Registration (called by the generators) ──────────────────────────────────
def count_records(module: str, out_bytes: bytes, sheet_count: int | None = None) -> Optional[int]:
    """Best-effort record count for the output workbook (None if it cannot be determined)."""
    if module == "loop_wiring":
        return sheet_count
    try:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(out_bytes), read_only=True, data_only=True)
        ws = wb.worksheets[-1]
        if module == "cable_schedule":
            # The output is a multi-sheet workbook; the cables live on the "Cable Schedule" sheet.
            ws = next((w for w in wb.worksheets if "cable schedule" in w.title.lower()), ws)
        rows = 0
        for row in ws.iter_rows(values_only=True):
            if any(c not in (None, "") for c in row):
                rows += 1
        wb.close()
        return max(rows - 1, 0)  # minus the header row
    except Exception:
        logger.warning("Could not count records for %s output", module, exc_info=True)
        return None


def _store_template(project_id: str, module: str, data: bytes, filename: str | None) -> str:
    tdir = _store_dir() / "projects" / project_id / module / "templates"
    tdir.mkdir(parents=True, exist_ok=True)
    path = tdir / f"{_sha256(data)[:12]}_{_safe_name(filename or 'template.xlsx')}"
    path.write_bytes(data)
    return str(path)


def _next_revision(db: Session, project_id: str, module: str) -> int:
    last = db.query(func.max(GeneratedDocument.revision)).filter(
        GeneratedDocument.project_id == project_id, GeneratedDocument.module == module,
        GeneratedDocument.generation_status == "success").scalar() or 0
    return last + 1


def latest_document(db: Session, project_id: str, module: str) -> Optional[GeneratedDocument]:
    """Latest *successful* revision of a module."""
    return (db.query(GeneratedDocument)
            .filter(GeneratedDocument.project_id == project_id, GeneratedDocument.module == module,
                    GeneratedDocument.generation_status == "success")
            .order_by(GeneratedDocument.revision.desc(), GeneratedDocument.id.desc()).first())


def register_generation(
    db: Session, *, project_id: str, module: str, user: str | None,
    source_bytes: bytes, source_filename: str,
    output_bytes: bytes, output_filename: str,
    params: dict | None = None, sheet_count: int | None = None, record_count: int | None = None,
    template_bytes: bytes | None = None, template_filename: str | None = None,
    project_name: str | None = None, action: str = "generated",
) -> GeneratedDocument:
    if module not in MODULES:
        raise RegistryError(f"Unknown module '{module}'.")
    pid = normalize_project_id(project_id)
    meta = MODULES[module]
    get_or_create_project(db, pid, project_name, user)
    sv = ensure_source_version(db, pid, meta["kind"], source_filename, source_bytes, user)

    prev = latest_document(db, pid, module)
    revision = _next_revision(db, pid, module)

    folder = _store_dir() / "projects" / pid / module
    folder.mkdir(parents=True, exist_ok=True)
    out_path = folder / f"rev{revision}_{_safe_name(output_filename)}"
    out_path.write_bytes(output_bytes)

    template_path = _store_template(pid, module, template_bytes, template_filename) if template_bytes else None

    if record_count is None:
        record_count = count_records(module, output_bytes, sheet_count)

    doc = GeneratedDocument(
        project_id=pid, module=module, document_name=meta["name"],
        document_number=f"{pid}-INS-{meta['serial']:03d}", revision=revision,
        status="GENERATED", generation_status="success",
        generated_by=user, source_version_id=sv.id, source_kind=sv.kind,
        source_rev_no=sv.rev_no, source_sha256=sv.sha256,
        filename=output_filename, output_file_path=str(out_path), template_path=template_path,
        record_count=record_count, sheet_count=sheet_count,
        params_json=json.dumps(params or {}),
        parent_document_id=prev.id if prev else None,
    )
    db.add(doc)
    db.flush()

    if prev:
        prev.status = "SUPERSEDED"
        prev.superseded_by = doc.id

    _audit(db, pid, action, user, module=module, document_id=doc.id,
           revision=revision, previous_revision=prev.revision if prev else None,
           source_kind=sv.kind, source_rev=sv.rev_no, filename=output_filename)
    db.commit()
    return doc


def register_failure(
    db: Session, *, project_id: str, module: str, user: str | None, error: str,
    source_bytes: bytes | None = None, source_filename: str | None = None,
    params: dict | None = None, template_path: str | None = None, project_name: str | None = None,
    template_bytes: bytes | None = None, template_filename: str | None = None,
) -> GeneratedDocument:
    """Record a failed generation. Never touches previously successful revisions."""
    pid = normalize_project_id(project_id)
    meta = MODULES[module]
    get_or_create_project(db, pid, project_name, user)
    if template_bytes and not template_path:
        template_path = _store_template(pid, module, template_bytes, template_filename)
    sv = None
    if source_bytes:
        sv = ensure_source_version(db, pid, meta["kind"], source_filename or "source.xlsx", source_bytes, user)
    doc = GeneratedDocument(
        project_id=pid, module=module, document_name=meta["name"],
        document_number=f"{pid}-INS-{meta['serial']:03d}", revision=0,
        status="FAILED", generation_status="failed", error_message=(error or "")[:2000],
        generated_by=user, source_version_id=sv.id if sv else None,
        source_kind=sv.kind if sv else meta["kind"], source_rev_no=sv.rev_no if sv else None,
        source_sha256=sv.sha256 if sv else None,
        params_json=json.dumps(params or {}), template_path=template_path,
    )
    db.add(doc)
    db.flush()
    _audit(db, pid, "generation_failed", user, module=module, document_id=doc.id, error=(error or "")[:500])
    db.commit()
    return doc


def register_safe(kind: str, **kwargs: Any) -> Optional[dict]:
    """
    Called from the generator routers. Opens its own session and never raises:
    a registry problem must not break a generation that already succeeded.

    `kind` is "success" or "failure". Returns None when no project was selected,
    otherwise {"registered": bool, "document": {...} | None, "error": str | None}
    so the UI can tell the engineer whether the register was updated.
    """
    if not kwargs.get("project_id"):
        return None
    from ..database import SessionLocal
    db = SessionLocal()
    try:
        if kind == "success":
            doc = register_generation(db, **kwargs)
            return {"registered": True, "document": doc_to_dict(doc), "error": None}
        register_failure(db, **kwargs)
        return {"registered": True, "document": None, "error": None}
    except RegistryError as exc:
        db.rollback()
        return {"registered": False, "document": None, "error": str(exc)}
    except Exception:
        db.rollback()
        logger.exception("Document registry registration failed (generation itself is unaffected)")
        return {"registered": False, "document": None, "error": "Document registry is unavailable"}
    finally:
        db.close()


# ── Read model ───────────────────────────────────────────────────────────────
def has_stored_settings(d: GeneratedDocument) -> bool:
    """Whether this row holds everything needed to re-run its generator."""
    if d.module == "loop_wiring":          # no parameters - only the template matters
        return bool(d.template_path)
    return bool(d.params_json and d.params_json != "{}")


def doc_to_dict(d: GeneratedDocument) -> dict:
    meta = MODULES.get(d.module, {})
    return {
        "id": d.id, "project_id": d.project_id, "module": d.module,
        "module_name": meta.get("name", d.module), "document_name": d.document_name,
        "document_number": d.document_number, "revision": d.revision,
        "revision_label": f"R{d.revision}" if d.revision else "—",
        "status": d.status, "generation_status": d.generation_status,
        "error_message": d.error_message,
        "generated_at": _iso(d.generated_at), "generated_by": d.generated_by,
        "source_kind": d.source_kind, "source_rev_no": d.source_rev_no,
        "source_sha256": d.source_sha256,
        "filename": d.filename, "record_count": d.record_count, "sheet_count": d.sheet_count,
        "count_label": meta.get("count_label", "Records"),
        "superseded_by": d.superseded_by, "parent_document_id": d.parent_document_id,
        "has_file": bool(d.output_file_path),
        "can_regenerate": has_stored_settings(d),
    }


def list_documents(db: Session, project_id: str, module: str | None = None) -> list[dict]:
    q = db.query(GeneratedDocument).filter(GeneratedDocument.project_id == project_id)
    if module:
        q = q.filter(GeneratedDocument.module == module)
    return [doc_to_dict(d) for d in q.order_by(GeneratedDocument.id.desc()).all()]


def active_link(db: Session, project_id: str, module: str) -> Optional[CoversheetLink]:
    return (db.query(CoversheetLink)
            .filter(CoversheetLink.project_id == project_id, CoversheetLink.module == module,
                    CoversheetLink.active.is_(True))
            .order_by(CoversheetLink.id.desc()).first())


def module_state(db: Session, project_id: str, module: str) -> dict:
    """Everything the Coversheet UI needs to know about one module."""
    meta = MODULES[module]
    latest = latest_document(db, project_id, module)
    last_attempt = (db.query(GeneratedDocument)
                    .filter(GeneratedDocument.project_id == project_id, GeneratedDocument.module == module)
                    .order_by(GeneratedDocument.id.desc()).first())
    link = active_link(db, project_id, module)
    linked_doc = db.get(GeneratedDocument, link.document_id) if link else None
    src = latest_source(db, project_id, meta["kind"])

    flags: list[str] = []
    if latest is None:
        state = "failed" if (last_attempt and last_attempt.generation_status == "failed") else "not_generated"
    else:
        if last_attempt and last_attempt.id > latest.id and last_attempt.generation_status == "failed":
            flags.append("last_generation_failed")
        if src and latest.source_rev_no is not None and latest.source_rev_no < src.rev_no:
            flags.append("regeneration_required")
        if link is None:
            flags.append("not_connected")
        elif link.document_id != latest.id:
            flags.append("new_revision_available")
        priority = ["last_generation_failed", "regeneration_required",
                    "new_revision_available", "not_connected"]
        state = next((f for f in priority if f in flags), "current")
        if state == "last_generation_failed":
            state = "failed"

    return {
        "module": module, "module_name": meta["name"], "source_kind": meta["kind"],
        "count_label": meta["count_label"],
        "state": state, "flags": flags,
        "connected": link is not None,
        "latest": doc_to_dict(latest) if latest else None,
        "linked": doc_to_dict(linked_doc) if linked_doc else None,
        "last_error": last_attempt.error_message
                      if last_attempt and last_attempt.generation_status == "failed"
                      and (latest is None or last_attempt.id > latest.id) else None,
        "latest_source_rev": src.rev_no if src else None,
    }


def overview(db: Session, project_id: str) -> dict:
    modules = [module_state(db, project_id, m) for m in MODULE_ORDER]
    iodb, loop = latest_source(db, project_id, "iodb"), latest_source(db, project_id, "loop_input")
    return {
        "project_id": project_id,
        "sources": {
            "iodb": {"rev_no": iodb.rev_no, "filename": iodb.filename, "created_at": _iso(iodb.created_at)} if iodb else None,
            "loop_input": {"rev_no": loop.rev_no, "filename": loop.filename, "created_at": _iso(loop.created_at)} if loop else None,
        },
        "modules": modules,
        "summary": summary_from(modules),
    }


def summary_from(modules: list[dict]) -> dict:
    def n(pred) -> int:
        return sum(1 for m in modules if pred(m))
    return {
        "total": len(modules),
        "current": n(lambda m: m["state"] == "current"),
        "regeneration_required": n(lambda m: "regeneration_required" in m["flags"]),
        "new_revision_available": n(lambda m: "new_revision_available" in m["flags"]),
        "not_connected": n(lambda m: m["state"] == "not_connected"),
        "not_generated": n(lambda m: m["state"] == "not_generated"),
        "failed": n(lambda m: m["state"] == "failed"),
        "coversheet_ready": all(m["state"] == "current" for m in modules if m["latest"] or m["connected"])
                            and any(m["connected"] for m in modules),
    }


# ── Connect / disconnect / status ────────────────────────────────────────────
def connect(db: Session, project_id: str, modules: list[str], user: str | None) -> dict:
    connected, skipped = [], []
    for module in modules:
        if module not in MODULES:
            raise RegistryError(f"Unknown module '{module}'.")
        latest = latest_document(db, project_id, module)
        if latest is None:
            skipped.append({"module": module, "reason": "Not generated yet"})
            continue
        current = active_link(db, project_id, module)
        if current and current.document_id == latest.id:
            connected.append(module)
            continue
        if current:
            current.active = False
            current.unlinked_by = user
            current.unlinked_at = datetime.utcnow()
        db.add(CoversheetLink(project_id=project_id, module=module, document_id=latest.id, linked_by=user))
        _audit(db, project_id, "connected", user, module=module, document_id=latest.id,
               revision=latest.revision,
               previous_linked_document=current.document_id if current else None)
        connected.append(module)
    db.commit()
    return {"connected": connected, "skipped": skipped}


def disconnect(db: Session, project_id: str, module: str, user: str | None) -> dict:
    if module not in MODULES:
        raise RegistryError(f"Unknown module '{module}'.")
    link = active_link(db, project_id, module)
    if not link:
        return {"disconnected": False, "message": "Module was not connected"}
    link.active = False
    link.unlinked_by = user
    link.unlinked_at = datetime.utcnow()
    _audit(db, project_id, "disconnected", user, module=module, document_id=link.document_id)
    db.commit()
    return {"disconnected": True, "message": "Disconnected. The generated file was not deleted."}


def set_status(db: Session, project_id: str, document_id: int, status: str, user: str | None) -> dict:
    status = (status or "").upper().replace(" ", "_")
    if status not in MANUAL_STATUSES:
        raise RegistryError(f"Status must be one of: {', '.join(MANUAL_STATUSES)}.")
    doc = db.get(GeneratedDocument, document_id)
    if not doc or doc.project_id != project_id:
        raise RegistryError("Document not found in this project.")
    if doc.status in ("SUPERSEDED", "FAILED"):
        raise RegistryError(f"A {doc.status.lower()} revision cannot change status.")
    old = doc.status
    doc.status = status
    _audit(db, project_id, "status_changed", user, module=doc.module, document_id=doc.id,
           old=old, new=status)
    db.commit()
    return doc_to_dict(doc)


def audit_trail(db: Session, project_id: str, limit: int = 200) -> list[dict]:
    rows = (db.query(DocumentAuditLog).filter(DocumentAuditLog.project_id == project_id)
            .order_by(DocumentAuditLog.id.desc()).limit(limit).all())
    return [{
        "id": r.id, "action": r.action, "module": r.module, "document_id": r.document_id,
        "actor": r.actor, "at": _iso(r.at),
        "details": json.loads(r.details) if r.details else {},
    } for r in rows]


# ── Regeneration (re-runs the EXISTING generators, never reimplements them) ──
def _run_generator(module: str, src: bytes, tmpl: bytes | None, params: dict
                   ) -> tuple[bytes | None, str | None, str | None, int | None]:
    """Returns (output_bytes, filename, error, sheet_count)."""
    _ensure_utils_importable()
    if module in ("instrument_list", "io_list"):
        from utils.template_processor import process_instrument_list, process_io_list
        fn = process_instrument_list if module == "instrument_list" else process_io_list
        out, name, err = fn(
            io.BytesIO(src), params.get("selected_columns") or [],
            filters=params.get("filters") or None,
            template_file=io.BytesIO(tmpl) if tmpl else None,
        )
        return out, name, err, None
    if module == "cable_schedule":
        from utils.template_processor import process_cable_schedule
        if not tmpl:
            return None, None, "Cable Schedule template is not stored; generate it once from the Cable Schedule page.", None
        out, name, err = process_cable_schedule(
            io.BytesIO(src), io.BytesIO(tmpl),
            params.get("jb_column") or "JUNCTION BOX", params.get("tag_column") or "TAG NO",
        )
        return out, name, err, None
    if module == "loop_wiring":
        from utils.file_handler import read_loop_wiring_input, load_workbook_from_upload
        from utils.generators import generate_loop_wiring
        from ..routers.loop_wiring import _find_tag_column
        if not tmpl:
            return None, None, "Loop Wiring template is not stored; generate it once from the Loop Wiring page.", None
        df, err = read_loop_wiring_input(io.BytesIO(src))
        if err:
            return None, None, err, None
        tag_col = _find_tag_column(df)
        sheets = len(df[tag_col].dropna().astype(str).str.strip().unique()) if tag_col else 0
        wb, buf, err = load_workbook_from_upload(io.BytesIO(tmpl))
        if err:
            return None, None, f"Failed to load Loop Wiring template: {err}", None
        buf.seek(0)
        out, name, err = generate_loop_wiring(df, wb, template_bytes=buf.read())
        return out, name, err, sheets
    return None, None, f"Unknown module '{module}'", None


def regenerate(db: Session, project_id: str, document_id: int, user: str | None) -> dict:
    """Regenerate a module from the project's LATEST source revision, as a new revision."""
    base = db.get(GeneratedDocument, document_id)
    if not base or base.project_id != project_id:
        raise RegistryError("Document not found in this project.")
    return regenerate_module(db, project_id, base.module, user, base=base)


def _latest_with_params(db: Session, project_id: str, module: str) -> Optional[GeneratedDocument]:
    rows = (db.query(GeneratedDocument)
            .filter(GeneratedDocument.project_id == project_id, GeneratedDocument.module == module)
            .order_by(GeneratedDocument.id.desc()).all())
    return next((d for d in rows if has_stored_settings(d)), None)


def regenerate_module(db: Session, project_id: str, module: str, user: str | None,
                      base: GeneratedDocument | None = None) -> dict:
    base = base or _latest_with_params(db, project_id, module)
    if base is None or not has_stored_settings(base):
        raise RegistryError(
            f"{MODULES[module]['name']} has never been generated with stored settings. "
            "Generate it once from its own page first."
        )
    meta = MODULES[module]
    src = latest_source(db, project_id, meta["kind"])
    if src is None or not Path(src.file_path).exists():
        raise RegistryError("No stored source file for this project. Upload the IODB first.")

    params = json.loads(base.params_json or "{}")
    tmpl = Path(base.template_path).read_bytes() if base.template_path and Path(base.template_path).exists() else None
    src_bytes = Path(src.file_path).read_bytes()

    try:
        out, name, err, sheets = _run_generator(module, src_bytes, tmpl, params)
    except Exception as exc:  # generator crashed - record, don't lose earlier revisions
        logger.exception("Regeneration of %s failed", module)
        out, name, err, sheets = None, None, str(exc), None

    if err or not out:
        fail = register_failure(db, project_id=project_id, module=module, user=user,
                                error=err or "Generator returned no output",
                                params=params, template_path=base.template_path)
        return {"ok": False, "module": module, "error": fail.error_message, "document": doc_to_dict(fail)}

    doc = register_generation(
        db, project_id=project_id, module=module, user=user,
        source_bytes=src_bytes, source_filename=src.filename,
        output_bytes=out, output_filename=name or base.filename or "output.xlsx",
        params=params, sheet_count=sheets, template_bytes=tmpl,
        template_filename=Path(base.template_path).name if base.template_path else None,
        action="regenerated",
    )
    return {"ok": True, "module": module, "document": doc_to_dict(doc)}


def generate_connected(db: Session, project_id: str, modules: list[str], user: str | None) -> dict:
    """
    Sequentially regenerate the selected modules from the latest stored source.
    A failure never stops later modules and never removes earlier outputs.
    """
    ordered = [m for m in MODULE_ORDER if m in modules]
    steps: list[dict] = []

    src = latest_source(db, project_id, "iodb")
    if src is None or not Path(src.file_path).exists():
        steps.append({"step": "iodb_check", "label": "Validating IODB", "status": "failed",
                      "message": "No IODB stored for this project. Upload or generate from an IODB first."})
        return {"steps": steps, "all_ok": False, "coversheet_ready": False}
    try:
        _ensure_utils_importable()
        from utils.file_handler import read_iodb
        _df, err = read_iodb(io.BytesIO(Path(src.file_path).read_bytes()))
    except Exception as exc:
        err = str(exc)
    steps.append({"step": "iodb_check", "label": f"Validating IODB (Rev {src.rev_no})",
                  "status": "failed" if err else "ok", "message": err or "IODB readable"})
    if err:
        return {"steps": steps, "all_ok": False, "coversheet_ready": False}

    for module in ordered:
        label = f"Generating {MODULES[module]['name']}"
        try:
            res = regenerate_module(db, project_id, module, user)
            steps.append({"step": module, "label": label,
                          "status": "ok" if res["ok"] else "failed",
                          "message": res.get("error") or f"Revision R{res['document']['revision']} generated",
                          "document": res["document"]})
        except RegistryError as exc:
            steps.append({"step": module, "label": label, "status": "needs_input", "message": str(exc)})

    all_ok = all(s["status"] == "ok" for s in steps)
    return {"steps": steps, "all_ok": all_ok, "coversheet_ready": all_ok}


def document_file(db: Session, project_id: str, document_id: int) -> tuple[Path, str]:
    doc = db.get(GeneratedDocument, document_id)
    if not doc or doc.project_id != project_id:
        raise RegistryError("Document not found in this project.")
    if not doc.output_file_path or not Path(doc.output_file_path).exists():
        raise RegistryError("The generated file is no longer available on the server.")
    return Path(doc.output_file_path), doc.filename or Path(doc.output_file_path).name

"""
Cable summary: robust IODB field resolution, sheet / header-row detection and diagnostics.
Regression for "No cable type / length data found" being shown for a workbook whose cable types
ARE present (the message must say what is really wrong).
"""
import io
from pathlib import Path

import openpyxl
import pandas as pd
import pytest

from utils.cable_summary import (CABLE_SUMMARY_SHEET, append_cable_summary_sheet, build_cable_summary,
                                 formula_hint_for, inspect_workbook, locate_cable_source, resolve_cable_fields)

ROOT = Path(__file__).resolve().parents[2]
REAL = Path("/Users/adithyachoudhrym/Downloads/iEZ Inputs/Cable Schedule/IODB.xlsx")
HEAD = ["TAG NO", "JUNCTION BOX", "INSTRUMENT TO JB CABLE TYPE", "BRANCH CABLE LENGTH", "MAIN CABLE TYPE",
        "MAIN CABLE LENGTH(Mtrs)"]


def workbook(rows, headers=HEAD, sheet="IODB - Table 1", header_row=1, blank_after_header=True, extra_sheets=()):
    wb = openpyxl.Workbook()
    for name in extra_sheets:
        wb.active.title = name
        wb.active["A1"] = "not the data"
        wb.create_sheet("tmp")
    ws = wb.active if not extra_sheets else wb.create_sheet(sheet)
    if not extra_sheets:
        ws.title = sheet
    for r in range(1, header_row):
        ws.cell(r, 1, "Project title / notes above the table")
    for c, h in enumerate(headers, 1):
        ws.cell(header_row, c, h)
    r = header_row + (2 if blank_after_header else 1)
    for row in rows:
        for c, v in enumerate(row, 1):
            ws.cell(r, c, v)
        r += 1
    if "tmp" in wb.sheetnames:
        del wb["tmp"]
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ── canonical field resolution ──────────────────────────────────────────────
@pytest.mark.parametrize("header", [
    "MAIN CABLE LENGTH(Mtrs)", "MAIN CABLE LENGTH(MTRS)", "MAIN CABLE LENGTH (Mtrs)", "MAIN CABLE LENGTH (MTRS)",
    "main cable length(mtrs)", "MAIN CABLE LENGTH\n(Mtrs)", "  MAIN  CABLE  LENGTH ( Mtrs ) ", "MAIN CABLE LENGTH",
    "MAIN CABLE LENGTH (M)", "Main Cable Length in Meters",
])
def test_main_length_header_variants_resolve_to_one_field(header):
    assert resolve_cable_fields([header])["MAIN_CABLE_LENGTH"] == header


@pytest.mark.parametrize("header", ["MAIN CABLE LENGTH TOLERANCE", "MAIN CABLE LENGTHS PER JB", "SUB MAIN CABLE LENGTH",
                                    "MAIN CABLE TAG NUMBER"])
def test_unrelated_headers_are_not_matched(header):
    assert resolve_cable_fields([header])["MAIN_CABLE_LENGTH"] is None


def test_types_and_branch_length_resolve_with_spacing_and_case():
    r = resolve_cable_fields(["Instrument to JB Cable Type ", "\nBRANCH CABLE LENGTH ", "main cable type"])
    assert r["INSTRUMENT_TO_JB_CABLE_TYPE"] == "Instrument to JB Cable Type "
    assert r["BRANCH_CABLE_LENGTH"] == "\nBRANCH CABLE LENGTH " and r["MAIN_CABLE_TYPE"] == "main cable type"


def test_column_position_is_irrelevant():
    cols = ["x1", "MAIN CABLE LENGTH(Mtrs)", "y", "MAIN CABLE TYPE", "z", "BRANCH CABLE LENGTH", "INSTRUMENT TO JB CABLE TYPE"]
    df = pd.DataFrame([[0, 100, 0, "12PX1.5 Sqmm", 0, 50, "1PX1.5 Sqmm"]], columns=cols)
    assert [(r["cable_description"], r["quantity"]) for r in build_cable_summary(df).rows] == [
        ("1PX1.5 Sqmm", 50), ("12PX1.5 Sqmm", 100)]


# ── sheet and header-row detection ──────────────────────────────────────────
def test_sheet_name_and_blank_row_after_header_do_not_matter():
    data = workbook([("T1", "JB1", "1PX1.5 Sqmm", 50, "12PX1.5 Sqmm", 100), ("T2", "JB1", "1PX1.5 Sqmm", "50", "12PX1.5 Sqmm", " 100 ")])
    found = inspect_workbook(data)
    assert found[0]["sheet"] == "IODB - Table 1" and found[0]["header_row"] == 1
    assert found[0]["columns"] == {"INSTRUMENT_TO_JB_CABLE_TYPE": "C", "BRANCH_CABLE_LENGTH": "D",
                                   "MAIN_CABLE_TYPE": "E", "MAIN_CABLE_LENGTH": "F"}
    df, info = locate_cable_source(pd.DataFrame({"unrelated": [1]}), data, "TAG NO")   # frame without cable fields -> rescan the workbook
    r = build_cable_summary(df, "TAG NO")
    assert [(x["cable_description"], x["quantity"]) for x in r.rows] == [("1PX1.5 Sqmm", 100), ("12PX1.5 Sqmm", 200)]


@pytest.mark.parametrize("sheet", ["IODB", "IODB - Format", "IODB - Table 1", "Sheet1", "Anything at all"])
def test_any_sheet_name_is_found_by_its_headers(sheet):
    assert inspect_workbook(workbook([("T1", "JB1", "1PX1.5 Sqmm", 5, None, None)], sheet=sheet))[0]["sheet"] == sheet


def test_header_row_need_not_be_row_one():
    data = workbook([("T1", "JB1", "1PX1.5 Sqmm", 5, "12PX1.5 Sqmm", 7)], header_row=4)
    assert inspect_workbook(data)[0]["header_row"] == 4
    df, info = locate_cable_source(pd.DataFrame({"x": [1]}), data)
    assert [(r["cable_description"], r["quantity"]) for r in build_cable_summary(df).rows] == [
        ("1PX1.5 Sqmm", 5), ("12PX1.5 Sqmm", 7)]


def test_correct_sheet_is_chosen_when_another_sheet_comes_first():
    data = workbook([("T1", "JB1", "1PX1.5 Sqmm", 5, None, None)], extra_sheets=("IODB Summary",))
    found = inspect_workbook(data)
    assert [f["sheet"] for f in found] == ["IODB - Table 1"]


# ── values: placeholders, numeric text, the empty-sheet reason ─────────────
def test_placeholders_are_ignored_and_numeric_text_is_accepted():
    df = pd.DataFrame(HEAD and [("T1", "JB", "-", "-", "TBA", "TBA"), ("T2", "JB", "1PX1.5 Sqmm", "TBA", "12PX1.5 Sqmm", "N/A"),
                                ("T3", "JB", "1PX1.5 Sqmm", "50", "12PX1.5 Sqmm", 100), ("T4", "JB", "1PX1.5 Sqmm", 50.0, None, None)],
                      columns=HEAD)
    r = build_cable_summary(df, "TAG NO")
    assert [(x["cable_description"], x["quantity"]) for x in r.rows] == [("1PX1.5 Sqmm", 100), ("12PX1.5 Sqmm", 100)]
    assert r.typed == {"branch cable": 3, "main cable": 2} and r.counts["blank_length"] == 2


def test_types_present_but_lengths_tba_gets_a_specific_reason_not_no_data():
    df = pd.DataFrame([("T1", "JB", "1PX1.5 Sqmm", "TBA", "12PX1.5 Sqmm", "TBA"), ("T2", "JB", "-", "-", "-", "-")], columns=HEAD)
    r = build_cable_summary(df)
    reason = r.empty_reason()
    assert r.rows == [] and "2 cable entries have a cable type" in reason and "none has a usable numeric length" in reason
    assert "TBA" in reason and "No cable type / length data found" not in reason
    wb = openpyxl.Workbook()
    append_cable_summary_sheet(wb, r)
    text = "".join(str(c.value) for row in wb[CABLE_SUMMARY_SHEET].iter_rows() for c in row if c.value)
    assert "none has a usable numeric length" in text
    assert not any(c.value == 0 for row in wb[CABLE_SUMMARY_SHEET].iter_rows() for c in row)


def test_no_types_at_all_has_its_own_reason():
    df = pd.DataFrame([("T1", "JB", "-", "-", None, None)], columns=HEAD)
    assert build_cable_summary(df).empty_reason().startswith("No cable types found")


def test_missing_columns_reason_names_them():
    r = build_cable_summary(pd.DataFrame({"TAG NO": ["A"]}))
    assert r.empty_reason().startswith("Cable columns not found in the IODB") and "MAIN CABLE LENGTH" in r.empty_reason()


def test_optional_columns_such_as_route_number_are_not_required():
    df = pd.DataFrame([("T1", "1PX1.5 Sqmm", 10, "12PX1.5 Sqmm", 20)],
                      columns=["TAG NO", "INSTRUMENT TO JB CABLE TYPE", "BRANCH CABLE LENGTH", "MAIN CABLE TYPE", "MAIN CABLE LENGTH(Mtrs)"])
    assert [r["quantity"] for r in build_cable_summary(df).rows] == [10, 20]       # no ROUTE NUMBER / JUNCTION BOX needed


# ── formulas whose results were never stored ───────────────────────────────
def test_formula_lengths_without_stored_results_are_diagnosed():
    data = workbook([("T1", "JB1", "1PX1.5 Sqmm", "=25*2", "12PX1.5 Sqmm", "=50*2")])      # openpyxl stores no cached results
    df = pd.read_excel(io.BytesIO(data), sheet_name=0).dropna(how="all")
    src, info = locate_cable_source(df, data)
    r = build_cable_summary(src)
    assert r.rows == []
    hint = formula_hint_for(data, info)
    assert "formulas with no stored result" in hint and "save it" in hint


# ── diagnostics ────────────────────────────────────────────────────────────
def test_diagnostics_report_sheet_header_row_columns_and_counts():
    data = workbook([("T1", "JB1", "1PX1.5 Sqmm", 50, "12PX1.5 Sqmm", 100), ("T2", "JB1", "1PX1.5 Sqmm", 50, "12PX1.5 Sqmm", 100)])
    df = pd.read_excel(io.BytesIO(data), sheet_name=0).dropna(how="all")
    src, info = locate_cable_source(df, data)
    r = build_cable_summary(src)
    r.source = {"sheet": info["sheet"], "header_row": info["header_row"], "columns": info["columns"]}
    text = "\n".join(r.diagnostics())
    assert "IODB sheet: IODB - Table 1; header row: 1" in text
    assert "MAIN_CABLE_LENGTH -> 'MAIN CABLE LENGTH(Mtrs)' (column F)" in text
    assert "branch cable: 2 entries with a cable type, 2 with a valid length" in text


def test_log_row_numbers_match_excel_even_with_the_blank_row_after_the_header():
    data = workbook([("T1", "JB1", "1PX1.5 Sqmm", "TBA", None, None), ("T2", "JB1", "1PX1.5 Sqmm", "TBA", None, None)])
    df = pd.read_excel(io.BytesIO(data), sheet_name=0).dropna(how="all").reset_index(drop=True)   # what the shared reader hands over
    src, info = locate_cable_source(df, data, "TAG NO")
    details = " ".join(build_cable_summary(src, "TAG NO").details)
    assert "IODB row 3, tag T1" in details and "IODB row 4, tag T2" in details       # header=1, blank=2, data from 3


def test_empty_table_shows_one_specific_message_only():
    wb = openpyxl.Workbook()
    df = pd.DataFrame([("T1", "JB", "1PX1.5 Sqmm", "TBA", None, None)], columns=HEAD)
    append_cable_summary_sheet(wb, build_cable_summary(df))
    texts = [str(c.value) for row in wb[CABLE_SUMMARY_SHEET].iter_rows() for c in row if c.value]
    assert len([t for t in texts if "length" in t]) == 1 and not any(t.startswith("NOTE:") for t in texts)


# ── the real project IODB (skipped when not present) ───────────────────────
@pytest.mark.skipif(not REAL.exists(), reason="project IODB not available")
def test_real_project_iodb_as_uploaded_finds_types_and_explains_tba_lengths():
    raw = REAL.read_bytes()
    from utils.file_handler import read_iodb
    df, _ = read_iodb(io.BytesIO(raw))
    src, info = locate_cable_source(df, raw, "TAG NO")
    assert info["sheet"] == "IODB - Table 1" and info["header_row"] == 1
    assert info["columns"] == {"INSTRUMENT_TO_JB_CABLE_TYPE": "W", "BRANCH_CABLE_LENGTH": "X",
                               "MAIN_CABLE_TYPE": "AE", "MAIN_CABLE_LENGTH": "AF"}
    r = build_cable_summary(src, "TAG NO")
    assert r.typed == {"branch cable": 154, "main cable": 153}
    assert r.rows == [] and "307 cable entries have a cable type" in r.empty_reason()


@pytest.mark.skipif(not REAL.exists(), reason="project IODB not available")
def test_real_project_iodb_with_lengths_filled_in_totals_correctly():
    wb = openpyxl.load_workbook(REAL)
    ws = wb.worksheets[0]
    expected: dict[str, float] = {}
    for r in range(3, ws.max_row + 1):
        for t_col, l_col, value in ((23, 24, 50), (31, 32, "100")):                 # W/X branch, AE/AF main (numeric text for main)
            desc = ws.cell(r, t_col).value
            if desc not in (None, "-", "TBA"):
                ws.cell(r, l_col).value = value
                expected[desc] = expected.get(desc, 0) + float(value)
    buf = io.BytesIO()
    wb.save(buf)
    raw = buf.getvalue()
    from utils.file_handler import read_iodb
    df, _ = read_iodb(io.BytesIO(raw))
    r = build_cable_summary(locate_cable_source(df, raw, "TAG NO")[0], "TAG NO")
    got = {x["cable_description"]: x["quantity"] for x in r.rows}
    assert got == {k: (int(v) if v == int(v) else v) for k, v in expected.items()}
    assert set(got) == {"1PX1.5 Sqmm", "1PX2.5 Sqmm", "12PX1.5 Sqmm", "12PX2.5 Sqmm", "12X1.5 Sq.mm"}
    assert r.counts["branch cable"] == 154 and r.counts["main cable"] == 153
